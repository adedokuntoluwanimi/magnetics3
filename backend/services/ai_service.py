from __future__ import annotations

import csv
import io
import json
import logging
from functools import lru_cache
from pathlib import Path
import re
import zipfile
from datetime import datetime, timezone
from typing import Any

import pandas as pd

from backend.models import AuroraResponse
from backend.gcp.vertex_ai import ExportProviderError
from backend.logging_utils import log_event
from backend.services.preview_service import _auto_utm_zone, _utm_to_wgs84
from backend.services.processing_service import _infer_base_station_mask

logger = logging.getLogger(__name__)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


@lru_cache(maxsize=1)
def _cached_export_agent_text() -> str:
    try:
        root = Path(__file__).resolve().parents[2]
        spec_path = root / "export_agent.md"
        if not spec_path.exists():
            spec_path = Path("/app/export_agent.md")
        if not spec_path.exists():
            return ""
        return spec_path.read_text(encoding="utf-8", errors="replace")
    except Exception:
        return ""

_EXPORT_SYSTEM_PROMPT = """You are the GAIA Magnetics export authoring engine.
Generate client-ready DOCX, PDF, and PPTX packages from actual processed magnetic survey outputs.

Non-negotiable rules:
- use only the supplied project metadata, processing path, diagnostics, QA, generated layers, selected export options, and available figures
- do not invent geology, anomalies, structures, corrections, metrics, or recommendations
- do not include unsupported sections, placeholder strings, raw JSON, raw dict/list output, raw CSV rows, template filler, or repeated paragraphs
- keep scientific labels truthful: do not mislabel fallbacks as strict geophysical products
- reflect degraded QA or fallback-heavy runs honestly
- DOCX is detailed, PDF is polished and concise, PPTX is visual and condensed
- return compact JSON only
- keep strings concise and omit nonessential optional fields if space is limited
- never add prose or markdown outside the JSON object
- never begin a field that cannot be completed; prioritize valid finished JSON over richness

Return only valid JSON in the required schema. Do not add markdown fences or prose outside the JSON."""


class AIService:
    _MAX_REF_CHARS = 50000
    _MAX_CHAT_CONTEXT_CHARS = 24000
    _MAX_UPLOAD_SNIPPET_CHARS = 20000

    def __init__(self, store, chat_client, export_client, storage_backend=None) -> None:
        self._store = store
        self._chat_client = chat_client
        self._export_client = export_client
        self._storage = storage_backend

    def run_startup_checks(self) -> None:
        preflight = getattr(self._export_client, "preflight_model_availability", None)
        if callable(preflight):
            preflight()

    def generate_preview(self, project_id: str, task_id: str) -> AuroraResponse:
        return self.generate_response(project_id, task_id, location="preview")

    def generate_response(
        self,
        project_id: str,
        task_id: str,
        *,
        location: str,
        question: str | None = None,
        history: list[dict] | None = None,
        extra_results: dict | None = None,
        ui_context: dict | None = None,
    ) -> AuroraResponse:
        project = self._store.get_project(project_id)
        task = self._store.get_task(task_id)
        if not project or not task:
            raise ValueError("Project or task not found.")

        outputs = extra_results or self._load_full_results(task) or task.get("results", {}).get("data", {})
        if location == "preview":
            preview_outputs = self._build_preview_outputs(task)
            if preview_outputs:
                merged_outputs = dict(outputs or {})
                merged_outputs["points"] = preview_outputs.get("points") or merged_outputs.get("points") or []
                merged_validation = dict(merged_outputs.get("validation_summary") or {})
                merged_validation.update(preview_outputs.get("validation_summary") or {})
                merged_outputs["validation_summary"] = merged_validation
                if preview_outputs.get("stats"):
                    merged_outputs["stats"] = preview_outputs["stats"]
                outputs = merged_outputs
        stats = outputs.get("stats") or {}
        config = task.get("analysis_config") or {}
        system_prompt = self._build_chat_system_prompt(project, task, config, outputs, stats, location, ui_context or {})

        messages: list[dict[str, str]] = []
        for msg in (history or [])[-20:]:
            role = msg.get("role", "user")
            content = (msg.get("content") or "").strip()
            if role in {"user", "assistant"} and content:
                messages.append({"role": role, "content": content})
        messages.append({"role": "user", "content": question or "Provide the most helpful guidance for this screen."})

        try:
            text = self._chat_client.generate(system_prompt=system_prompt, messages=messages, max_tokens=2400)
        except Exception as exc:
            logger.exception("Aurora chat generation failed", extra={"location": location, "task_id": task_id, "project_id": project_id, "error": str(exc)})
            try:
                compact_prompt = self._build_compact_chat_system_prompt(project, task, config, outputs, stats, location, ui_context or {})
                text = self._chat_client.generate(system_prompt=compact_prompt, messages=messages[-6:], max_tokens=800)
            except Exception:
                text = self._fallback_text(project, task, outputs, location, question)
        return self._parse_chat_response(text, location)

    def generate_export_report(self, project_id: str, task_id: str, full_results: dict, export_context: dict | None = None) -> AuroraResponse:
        project = self._store.get_project(project_id)
        task = self._store.get_task(task_id)
        if not project or not task:
            raise ValueError("Project or task not found.")

        config = task.get("analysis_config") or {}
        stats = full_results.get("stats") or {}
        upload_context = self._collect_uploaded_file_context(task)
        export_context = export_context or {}
        export_type = ",".join(export_context.get("formats") or []) or "unspecified"
        provider = "anthropic"
        model = getattr(getattr(self._export_client, "_settings", None), "aurora_export_model", "unknown")
        provider_errors: list[str] = []
        fallback_package = self._build_fallback_export_package(project, task, config, full_results, stats, upload_context, export_context)
        block_package, block_status = self._generate_export_blocks(
            project=project,
            task=task,
            config=config,
            full_results=full_results,
            stats=stats,
            upload_context=upload_context,
            export_context=export_context,
            fallback_package=fallback_package,
            project_id=project_id,
            task_id=task_id,
            provider=provider,
            model=model,
        )
        provider_errors.extend(block_status.get("provider_errors") or [])
        final_outcome = "anthropic_success"
        if block_status.get("request_failed"):
            final_outcome = "anthropic_request_failed_fallback_used"
        if block_status.get("response_invalid") or block_status.get("validation_failed"):
            final_outcome = "anthropic_response_invalid_fallback_used"

        report_data = self._prepare_export_package(
            block_package,
            full_results,
            export_context,
            task_id=task_id,
            export_type=export_type,
            provider=provider,
            model=model,
            attempt_number=0,
        )
        if report_data and final_outcome == "anthropic_success":
            self._log_export_outcome(
                task_id=task_id,
                export_type=export_type,
                provider=provider,
                model=model,
                outcome="anthropic_success",
            )
            highlights = []
            primary_doc = report_data.get("pdf") or report_data.get("docx") or report_data.get("pptx") or {}
            for section in (primary_doc.get("sections") or [])[:5]:
                title = section.get("title") or section.get("type") or "Section"
                body = " ".join(section.get("bullets") or [])[:140]
                if body:
                    highlights.append(f"{title}: {body}")
            summary = primary_doc.get("executive_summary") or report_data.get("summary") or ""
            return AuroraResponse(
                location="export",
                summary=summary,
                highlights=highlights,
                message=summary,
                report_data=report_data,
            )

        fallback_context = dict(export_context)
        if provider_errors:
            fallback_context["provider_errors"] = provider_errors
        fallback_report = self._build_fallback_export_package(project, task, config, full_results, stats, upload_context, fallback_context)
        fallback_report = self._prepare_export_package(fallback_report, full_results, fallback_context) or fallback_report
        if final_outcome == "anthropic_success":
            final_outcome = "anthropic_response_invalid_fallback_used"
        self._log_export_outcome(
            task_id=task_id,
            export_type=export_type,
            provider=provider,
            model=model,
            outcome=final_outcome,
        )
        fallback_summary = fallback_report.get("summary", "")
        return AuroraResponse(
            location="export",
            summary=fallback_summary,
            highlights=(fallback_report.get("pdf") or {}).get("recommendations", [])[:4],
            message=fallback_summary,
            report_data=fallback_report,
        )

    def _generate_export_blocks(
        self,
        *,
        project: dict,
        task: dict,
        config: dict,
        full_results: dict,
        stats: dict,
        upload_context: str | None,
        export_context: dict[str, Any],
        fallback_package: dict[str, Any],
        project_id: str,
        task_id: str,
        provider: str,
        model: str,
    ) -> tuple[dict[str, Any], dict[str, Any]]:
        package = json.loads(json.dumps(fallback_package, default=str))
        status = {
            "provider_errors": [],
            "request_failed": False,
            "response_invalid": False,
            "validation_failed": False,
        }
        for block in self._export_block_definitions(full_results, export_context):
            result = self._run_export_block(
                project=project,
                task=task,
                config=config,
                full_results=full_results,
                stats=stats,
                upload_context=upload_context,
                export_context=export_context,
                fallback_package=fallback_package,
                project_id=project_id,
                task_id=task_id,
                provider=provider,
                model=model,
                block=block,
                current_package=package,
            )
            status["provider_errors"].extend(result.get("provider_errors") or [])
            status["request_failed"] = status["request_failed"] or bool(result.get("request_failed"))
            status["response_invalid"] = status["response_invalid"] or bool(result.get("response_invalid"))
            status["validation_failed"] = status["validation_failed"] or bool(result.get("validation_failed"))
            block_payload = result.get("payload")
            if isinstance(block_payload, dict):
                package = self._merge_export_block_payload(package, block, block_payload)
        return package, status

    def _export_block_definitions(self, full_results: dict, export_context: dict[str, Any]) -> list[dict[str, Any]]:
        blocks: list[dict[str, Any]] = []
        requested_keys = self._requested_export_document_keys(export_context)
        layer_titles = self._selected_export_layer_labels(full_results, dict(export_context.get("export_options") or {}))
        selected_sections = set(export_context.get("aurora_sections") or (export_context.get("export_options") or {}).get("selected_narrative_sections") or [])
        if {"pdf", "docx"} & set(requested_keys):
            blocks.extend(
                [
                    {"name": "executive_summary", "doc_key": "report", "target_titles": [], "max_tokens": 480, "retry_max_tokens": 320},
                    {"name": "project_setup", "doc_key": "report", "target_titles": ["Project Overview"], "max_tokens": 800, "retry_max_tokens": 520},
                    {"name": "task_summary", "doc_key": "report", "target_titles": ["Data and Survey Summary"], "max_tokens": 640, "retry_max_tokens": 420},
                    {"name": "processing_qa", "doc_key": "report", "target_titles": ["Processing Workflow Summary"], "max_tokens": 640, "retry_max_tokens": 420},
                    {"name": "key_findings", "doc_key": "report", "target_titles": ["Key Findings"], "max_tokens": 560, "retry_max_tokens": 380},
                    {"name": "result_layers", "doc_key": "report", "target_titles": layer_titles[:5], "max_tokens": 900, "retry_max_tokens": 580},
                    {"name": "limitations", "doc_key": "report", "target_titles": ["Data Quality and Reliability"], "max_tokens": 560, "retry_max_tokens": 380},
                    {"name": "recommendations", "doc_key": "report", "target_titles": ["Conclusions"], "max_tokens": 480, "retry_max_tokens": 320},
                ]
            )
        if "pptx" in requested_keys:
            pptx_layers = layer_titles[:3]
            blocks.extend(
                [
                    {"name": "pptx_group_1", "doc_key": "pptx", "target_titles": ["Executive Summary"], "max_tokens": 640, "retry_max_tokens": 400},
                    {"name": "pptx_group_2a", "doc_key": "pptx", "target_titles": ["Project Overview"], "max_tokens": 480, "retry_max_tokens": 320},
                    {"name": "pptx_group_2b", "doc_key": "pptx", "target_titles": ["Data and Survey Summary", "Processing Workflow Summary"], "max_tokens": 640, "retry_max_tokens": 420},
                    {"name": "pptx_group_3", "doc_key": "pptx", "target_titles": ["Key Findings"] + pptx_layers, "max_tokens": 760, "retry_max_tokens": 480},
                    {"name": "pptx_group_4", "doc_key": "pptx", "target_titles": ["Data Quality and Reliability", "Conclusions"], "max_tokens": 480, "retry_max_tokens": 320},
                ]
            )
        filtered: list[dict[str, Any]] = []
        for block in blocks:
            if block["name"] == "result_layers" and not block["target_titles"]:
                continue
            if block["name"] == "recommendations" and "Drilling recommendations" not in selected_sections and "Coverage gap analysis" not in selected_sections and "QA summary" not in selected_sections:
                block = dict(block)
            filtered.append(block)
        return filtered

    def _run_export_block(
        self,
        *,
        project: dict,
        task: dict,
        config: dict,
        full_results: dict,
        stats: dict,
        upload_context: str | None,
        export_context: dict[str, Any],
        fallback_package: dict[str, Any],
        project_id: str,
        task_id: str,
        provider: str,
        model: str,
        block: dict[str, Any],
        current_package: dict[str, Any],
    ) -> dict[str, Any]:
        provider_errors: list[str] = []
        result = {"payload": None, "provider_errors": provider_errors, "request_failed": False, "response_invalid": False, "validation_failed": False}
        prompt_context = self._build_export_block_context(
            project,
            task,
            config,
            full_results,
            stats,
            upload_context,
            export_context,
            fallback_package,
            block,
        )
        user_prompt = self._build_export_block_prompt(block, prompt_context, retry_mode="initial")
        self._log_export_token_estimate(task_id, project_id, _EXPORT_SYSTEM_PROMPT, user_prompt)
        try:
            text = self._export_client.generate(
                system_prompt=_EXPORT_SYSTEM_PROMPT,
                user_prompt=user_prompt,
                max_tokens=int(block.get("max_tokens") or 320),
            )
            parsed, parse_meta = self._parse_export_block_json(text)
            if parsed:
                validated, issues = self._validate_export_block_payload(block, parsed, full_results, export_context, current_package)
                if validated:
                    self._log_export_block_event(
                        task_id=task_id,
                        export_type=",".join(export_context.get("formats") or []) or "unspecified",
                        provider=provider,
                        model=model,
                        block_name=block["name"],
                        action="export.block.outcome",
                        outcome="success",
                        attempt_number=1,
                    )
                    result["payload"] = validated
                    return result
                self._log_export_failure(
                    task_id=task_id,
                    export_type=",".join(export_context.get("formats") or []) or "unspecified",
                    provider=provider,
                    model=model,
                    attempt_number=1,
                    phase="block_validation",
                    category="validation_rejected",
                    message="Export block was rejected during validation.",
                    validation_reasons=issues[:10],
                    block_name=block["name"],
                )
                result["validation_failed"] = True
                return result
            self._log_export_parse_forensics(
                task_id=task_id,
                export_type=",".join(export_context.get("formats") or []) or "unspecified",
                provider=provider,
                model=model,
                attempt_number=1,
                phase="api_response_parse",
                forensics=parse_meta.get("forensics") or [],
                block_name=block["name"],
            )
            self._log_export_failure(
                task_id=task_id,
                export_type=",".join(export_context.get("formats") or []) or "unspecified",
                provider=provider,
                model=model,
                attempt_number=1,
                phase="api_response_parse",
                category=parse_meta.get("failure_category") or "response_parse_failed",
                message="Export block response could not be parsed.",
                parse_strategy_attempted=parse_meta.get("strategies"),
                fenced_json_detected=parse_meta.get("fenced_json_detected"),
                object_extraction_attempted=parse_meta.get("object_extraction_attempted"),
                object_candidate_count=parse_meta.get("object_candidate_count"),
                response_preview=self._safe_preview(text),
                block_name=block["name"],
            )
            failure_class = self._primary_failure_class(parse_meta)
            retry_mode = "same_request_retry"
            compact_context = prompt_context
            if failure_class == "truncated_json":
                retry_mode = "compact_schema_retry"
                compact_context = self._compact_block_context(prompt_context)
            self._log_export_retry_mode(
                task_id=task_id,
                export_type=",".join(export_context.get("formats") or []) or "unspecified",
                provider=provider,
                model=model,
                retry_mode=retry_mode,
                attempt_number=2,
                reason=failure_class or "response_parse_failed",
                block_name=block["name"],
            )
        except ExportProviderError as exc:
            provider_errors.append(str(exc))
            self._log_export_failure(
                task_id=task_id,
                export_type=",".join(export_context.get("formats") or []) or "unspecified",
                provider=exc.provider,
                model=exc.model,
                attempt_number=1,
                phase="api_request",
                category=exc.category,
                message=exc.provider_message,
                request_id=exc.request_id,
                status_code=exc.status_code,
                retryable=exc.retryable,
                block_name=block["name"],
            )
            result["request_failed"] = True
            return result
        except Exception as exc:
            self._log_export_failure(
                task_id=task_id,
                export_type=",".join(export_context.get("formats") or []) or "unspecified",
                provider=provider,
                model=model,
                attempt_number=1,
                phase="api_request",
                category="unexpected_request_failure",
                message=str(exc),
                block_name=block["name"],
            )
            result["request_failed"] = True
            return result

        try:
            retry_prompt = self._build_export_block_prompt(block, compact_context, retry_mode=retry_mode)
            text = self._export_client.generate(
                system_prompt=_EXPORT_SYSTEM_PROMPT,
                user_prompt=retry_prompt,
                max_tokens=int(block.get("retry_max_tokens") or 220),
            )
            parsed, parse_meta = self._parse_export_block_json(text)
            if parsed:
                validated, issues = self._validate_export_block_payload(block, parsed, full_results, export_context, current_package)
                if validated:
                    self._log_export_block_event(
                        task_id=task_id,
                        export_type=",".join(export_context.get("formats") or []) or "unspecified",
                        provider=provider,
                        model=model,
                        block_name=block["name"],
                        action="export.block.outcome",
                        outcome="success_after_retry",
                        attempt_number=2,
                    )
                    result["payload"] = validated
                    return result
                self._log_export_failure(
                    task_id=task_id,
                    export_type=",".join(export_context.get("formats") or []) or "unspecified",
                    provider=provider,
                    model=model,
                    attempt_number=2,
                    phase="block_validation",
                    category="validation_rejected",
                    message="Export block retry was rejected during validation.",
                    validation_reasons=issues[:10],
                    block_name=block["name"],
                )
                result["validation_failed"] = True
                return result
            self._log_export_parse_forensics(
                task_id=task_id,
                export_type=",".join(export_context.get("formats") or []) or "unspecified",
                provider=provider,
                model=model,
                attempt_number=2,
                phase="api_response_parse",
                forensics=parse_meta.get("forensics") or [],
                block_name=block["name"],
            )
            self._log_export_failure(
                task_id=task_id,
                export_type=",".join(export_context.get("formats") or []) or "unspecified",
                provider=provider,
                model=model,
                attempt_number=2,
                phase="api_response_parse",
                category=parse_meta.get("failure_category") or "response_parse_failed",
                message="Export block retry response could not be parsed.",
                parse_strategy_attempted=parse_meta.get("strategies"),
                fenced_json_detected=parse_meta.get("fenced_json_detected"),
                object_extraction_attempted=parse_meta.get("object_extraction_attempted"),
                object_candidate_count=parse_meta.get("object_candidate_count"),
                response_preview=self._safe_preview(text),
                block_name=block["name"],
            )
            result["response_invalid"] = True
            self._log_export_block_event(
                task_id=task_id,
                export_type=",".join(export_context.get("formats") or []) or "unspecified",
                provider=provider,
                model=model,
                block_name=block["name"],
                action="export.block.outcome",
                outcome="fallback_used",
                attempt_number=2,
            )
            return result
        except ExportProviderError as exc:
            provider_errors.append(str(exc))
            self._log_export_failure(
                task_id=task_id,
                export_type=",".join(export_context.get("formats") or []) or "unspecified",
                provider=exc.provider,
                model=exc.model,
                attempt_number=2,
                phase="api_request",
                category=exc.category,
                message=exc.provider_message,
                request_id=exc.request_id,
                status_code=exc.status_code,
                retryable=exc.retryable,
                block_name=block["name"],
            )
            result["request_failed"] = True
            return result
        except Exception as exc:
            self._log_export_failure(
                task_id=task_id,
                export_type=",".join(export_context.get("formats") or []) or "unspecified",
                provider=provider,
                model=model,
                attempt_number=2,
                phase="api_request",
                category="unexpected_retry_failure",
                message=str(exc),
                block_name=block["name"],
            )
            result["request_failed"] = True
            return result

    def _build_export_block_context(
        self,
        project: dict,
        task: dict,
        config: dict,
        full_results: dict,
        stats: dict,
        upload_context: str | None,
        export_context: dict[str, Any],
        fallback_package: dict[str, Any],
        block: dict[str, Any],
    ) -> dict[str, Any]:
        block_name = block["name"]
        export_options = dict(export_context.get("export_options") or {})
        fallback_sections = self._fallback_sections_by_title(fallback_package, "pptx" if block.get("doc_key") == "pptx" else "pdf")
        selected_layers = self._selected_export_layer_labels(full_results, export_options)
        common = {
            "project": {
                "name": project.get("name", ""),
                "task_name": task.get("name", ""),
                "platform": task.get("platform", "not specified"),
                "scenario": task.get("scenario") or "automatic",
            },
            "qa_summary": self._build_compact_qa_summary(full_results),
            "processing_path": self._build_compact_processing_path(task, full_results, config.get("corrections") or [], config.get("add_ons") or []),
            "diagnostics_summary": self._build_compact_diagnostics_summary(full_results, stats, full_results.get("model_metadata") or {}),
            "selected_layers": selected_layers,
            "selected_visual_refs": [item["visual_ref"] for item in self._build_export_visual_refs(full_results)],
            "rtp_label": self._rtp_label(full_results),
        }
        if block_name == "executive_summary":
            common["fallback_block"] = {
                "title": (fallback_package.get("pdf") or {}).get("title"),
                "subtitle": (fallback_package.get("pdf") or {}).get("subtitle"),
                "executive_summary": (fallback_package.get("pdf") or {}).get("executive_summary"),
            }
        else:
            common["fallback_sections"] = [fallback_sections.get(title) for title in block.get("target_titles") or [] if fallback_sections.get(title)]
            if block_name in {"recommendations", "pptx_group_4"}:
                common["fallback_recommendations"] = (fallback_package.get("pdf") or {}).get("recommendations") or []
            if block_name in {"result_layers", "pptx_group_3"}:
                common["layer_summaries"] = self._build_compact_layer_summaries(full_results, block.get("target_titles") or [])
            if block_name in {"project_setup", "task_summary"} and upload_context:
                common["uploaded_files_summary"] = self._compact_text(upload_context, 220)
        return self._sanitize_export_context_payload(common)

    def _build_export_block_prompt(self, block: dict[str, Any], context: dict[str, Any], *, retry_mode: str) -> str:
        block_name = block["name"]
        doc_key = block.get("doc_key")
        limits = self._block_prompt_limits(block, retry_mode)
        if block_name == "executive_summary":
            schema = '- keys: title, subtitle, executive_summary'
        elif doc_key == "pptx":
            schema = '- keys: sections, recommendations\n- each section keys: type, title, bullets, visual_refs, speaker_notes, callouts'
        else:
            schema = '- keys: sections, recommendations\n- each section keys: type, title, bullets, body, visual_refs, interpretation, implication, callouts'
        return (
            f"Task: write the compact export block `{block_name}` for `{doc_key}`.\n"
            f"Retry mode: {retry_mode}\n"
            "Return one small JSON object only.\n"
            "Keep every string concise.\n"
            "Use only the provided context and fallback block content.\n"
            "Do not repeat text already implied by QA or processing path.\n"
            "Do not output markdown, prose, or comments.\n"
            "If space is tight, shorten bullets before dropping required meaning.\n"
            "Keep scientific labels honest.\n"
            f"Max sections: {limits['max_sections']}. Max bullets per section: {limits['max_bullets']}. Max bullet chars: {limits['max_bullet_chars']}.\n"
            f"Max body chars: {limits['max_body_chars']}. Max interpretation chars: {limits['max_interpretation_chars']}. Max recommendation chars: {limits['max_recommendation_chars']}.\n"
            "Do not wrap the payload inside the block name.\n"
            + schema
            + "\nContext JSON:\n"
            + json.dumps(context, separators=(",", ":"), ensure_ascii=True)
        )

    def _block_prompt_limits(self, block: dict[str, Any], retry_mode: str) -> dict[str, int]:
        base = {
            "max_sections": 2 if block.get("doc_key") == "pptx" else 3,
            "max_bullets": 3 if block.get("doc_key") == "pptx" else 3,
            "max_bullet_chars": 70 if block.get("doc_key") == "pptx" else 95,
            "max_body_chars": 140 if block.get("doc_key") == "pptx" else 220,
            "max_interpretation_chars": 100 if block.get("doc_key") == "pptx" else 140,
            "max_recommendation_chars": 110,
        }
        if block["name"] == "executive_summary":
            base.update({"max_sections": 0, "max_bullets": 0, "max_body_chars": 0, "max_interpretation_chars": 0})
        if block["name"] in {"pptx_group_4", "recommendations"}:
            base.update({"max_sections": 2, "max_bullets": 2, "max_bullet_chars": 60, "max_body_chars": 110, "max_interpretation_chars": 90, "max_recommendation_chars": 95})
        if retry_mode == "compact_schema_retry":
            base["max_sections"] = min(base["max_sections"], 2)
            base["max_bullets"] = min(base["max_bullets"], 2)
            base["max_bullet_chars"] = min(base["max_bullet_chars"], 55 if block.get("doc_key") == "pptx" else 80)
            base["max_body_chars"] = min(base["max_body_chars"], 90 if block.get("doc_key") == "pptx" else 160)
            base["max_interpretation_chars"] = min(base["max_interpretation_chars"], 70 if block.get("doc_key") == "pptx" else 110)
            base["max_recommendation_chars"] = min(base["max_recommendation_chars"], 85)
        return base

    def _compact_block_context(self, context: dict[str, Any]) -> dict[str, Any]:
        compact = json.loads(json.dumps(context, default=str))
        if isinstance(compact.get("fallback_sections"), list):
            compact["fallback_sections"] = compact["fallback_sections"][:2]
            for section in compact["fallback_sections"]:
                if not isinstance(section, dict):
                    continue
                if isinstance(section.get("bullets"), list):
                    section["bullets"] = [self._compact_text(item, 90) for item in section["bullets"][:3]]
                section["body"] = self._compact_text(section.get("body"), 160)
                section["interpretation"] = self._compact_text(section.get("interpretation"), 120)
                section["implication"] = self._compact_text(section.get("implication"), 120)
        if isinstance(compact.get("layer_summaries"), list):
            compact["layer_summaries"] = compact["layer_summaries"][:3]
        if isinstance(compact.get("fallback_recommendations"), list):
            compact["fallback_recommendations"] = [self._compact_text(item, 110) for item in compact["fallback_recommendations"][:2]]
        if isinstance(compact.get("qa_summary"), dict):
            compact["qa_summary"]["warnings"] = [self._compact_text(item, 80) for item in (compact["qa_summary"].get("warnings") or [])[:2]]
            compact["qa_summary"]["fallback_events"] = (compact["qa_summary"].get("fallback_events") or [])[:2]
        return compact

    def _fallback_sections_by_title(self, package: dict[str, Any], key: str) -> dict[str, dict[str, Any]]:
        payload = package.get(key) or {}
        return {str(section.get("title") or ""): json.loads(json.dumps(section, default=str)) for section in (payload.get("sections") or []) if section.get("title")}

    def _build_compact_layer_summaries(self, full_results: dict, target_titles: list[str]) -> list[dict[str, Any]]:
        summaries: list[dict[str, Any]] = []
        for title in target_titles:
            grid = full_results.get(self._result_key_for_export_layer(title))
            stats = self._grid_stats(grid)
            entry = {"title": title, "visual_refs": [item["visual_ref"] for item in self._build_export_visual_refs(full_results) if item["layer"] == title][:2]}
            if stats:
                entry["range_nT"] = round(float(stats.get("range", 0) or 0), 2)
                entry["std_nT"] = round(float(stats.get("std", 0) or 0), 2)
            summaries.append(entry)
        return summaries

    def _result_key_for_export_layer(self, title: str) -> str:
        mapping = {
            "Corrected Magnetic Field": "surface",
            "Regional Magnetic Field": "regional_field",
            "Residual Magnetic Field": "regional_residual",
            "Reduction to Pole (RTP)": "rtp_surface",
            "Low-latitude fallback transform": "rtp_surface",
            "Analytic Signal": "analytic_signal",
            "Tilt Derivative": "tilt_derivative",
            "Total Gradient": "total_gradient",
            "First Vertical Derivative": "first_vertical_derivative",
            "First Vertical Derivative (FVD)": "first_vertical_derivative",
            "Horizontal Derivative": "horizontal_derivative",
            "Uncertainty": "uncertainty",
            "Line Profiles": "line_profile_b64",
        }
        return mapping.get(title, "")

    def _parse_export_block_json(self, text: str) -> tuple[dict | None, dict[str, Any]]:
        parsed, meta = self._parse_json_object(text)
        if parsed is None:
            return None, meta
        if not isinstance(parsed, dict):
            meta["failure_category"] = "response_parse_failed"
            return None, meta
        return parsed, meta

    def _unwrap_block_wrapper(self, block: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any]:
        lowered = block["name"].lower()
        if lowered in {str(key).lower() for key in payload.keys()} and isinstance(payload.get(block["name"]), dict):
            return payload[block["name"]]
        if len(payload) == 1:
            only_value = next(iter(payload.values()))
            if isinstance(only_value, dict):
                return only_value
        return payload

    def _validate_export_block_payload(
        self,
        block: dict[str, Any],
        payload: dict[str, Any],
        full_results: dict,
        export_context: dict[str, Any],
        current_package: dict[str, Any],
    ) -> tuple[dict[str, Any] | None, list[str]]:
        issues: list[str] = []
        cleaned = self._unwrap_block_wrapper(block, json.loads(json.dumps(payload, default=str)))
        if block["name"] == "executive_summary":
            for key in ("title", "subtitle", "executive_summary"):
                if key in cleaned:
                    cleaned[key] = self._compact_text(cleaned.get(key), 220 if key != "executive_summary" else 420)
            if not cleaned.get("executive_summary"):
                issues.append("missing_executive_summary")
        else:
            sections = []
            for section in (cleaned.get("sections") or []):
                if not isinstance(section, dict):
                    continue
                title = str(section.get("title") or "").strip()
                section["title"] = title
                section["type"] = str(section.get("type") or "section").strip()
                section["include"] = True
                section["bullets"] = [self._compact_text(item, 140 if block.get("doc_key") != "pptx" else 90) for item in (section.get("bullets") or []) if str(item).strip()][: (5 if block.get("doc_key") != "pptx" else 4)]
                section["body"] = self._compact_text(section.get("body"), 320 if block.get("doc_key") != "pptx" else 180)
                section["interpretation"] = self._compact_text(section.get("interpretation"), 180 if block.get("doc_key") != "pptx" else 120)
                section["implication"] = self._compact_text(section.get("implication"), 180 if block.get("doc_key") != "pptx" else 120)
                section["caption"] = self._compact_text(section.get("caption"), 120)
                section["speaker_notes"] = self._compact_text(section.get("speaker_notes"), 180)
                section["callouts"] = [self._compact_text(item, 120) for item in (section.get("callouts") or []) if str(item).strip()][:3]
                section["visual_refs"] = [str(item).strip() for item in (section.get("visual_refs") or []) if str(item).strip()][:2]
                sections.append(section)
            cleaned["sections"] = sections
            cleaned["recommendations"] = [self._compact_text(item, 150) for item in (cleaned.get("recommendations") or []) if str(item).strip()][:4]
            expected_titles = set(block.get("target_titles") or [])
            normalized_expected = {self._normalize_block_title(title) for title in expected_titles}
            if expected_titles and not any(self._normalize_block_title(section.get("title") or "") in normalized_expected for section in sections):
                issues.append("missing_expected_section")
            existing_texts = self._collect_existing_export_texts(current_package, block.get("doc_key"))
            for section in sections:
                text_blob = " ".join(section.get("bullets") or []) + " " + (section.get("body") or "") + " " + (section.get("interpretation") or "")
                normalized = re.sub(r"\s+", " ", text_blob).strip().lower()
                if normalized and normalized in existing_texts:
                    issues.append("repeated_block_content")
                if any(fragment in normalized for fragment in ("latitude,longitude", "\"type\":", "[insert", "todo", "lorem ipsum")):
                    issues.append("blocked_content")
            if cleaned.get("recommendations"):
                low = " ".join(cleaned["recommendations"]).lower()
                if not any(term in low for term in ("support", "uncertainty", "fallback", "level", "base-station", "prediction", "inclination", "declination", "crossover", "qa")):
                    issues.append("untethered_recommendations")
        return (cleaned if not issues else None), issues

    def _normalize_block_title(self, value: str) -> str:
        return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())

    def _collect_existing_export_texts(self, package: dict[str, Any], doc_key: str | None) -> set[str]:
        keys = ["pptx"] if doc_key == "pptx" else ["pdf", "docx"]
        seen: set[str] = set()
        for key in keys:
            payload = package.get(key) or {}
            values = [payload.get("executive_summary") or ""] + [str(item) for item in (payload.get("recommendations") or [])]
            for section in payload.get("sections") or []:
                values.append(" ".join(section.get("bullets") or []))
                values.append(section.get("body") or "")
                values.append(section.get("interpretation") or "")
            for value in values:
                normalized = re.sub(r"\s+", " ", str(value).strip()).lower()
                if normalized:
                    seen.add(normalized)
        return seen

    def _merge_export_block_payload(self, package: dict[str, Any], block: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any]:
        merged = json.loads(json.dumps(package, default=str))
        if block["doc_key"] == "report":
            for key in ("pdf", "docx"):
                merged[key] = self._apply_block_payload_to_document(merged.get(key) or {}, block, payload)
            if payload.get("executive_summary"):
                merged["summary"] = payload.get("executive_summary")
        else:
            merged["pptx"] = self._apply_block_payload_to_document(merged.get("pptx") or {}, block, payload)
        return merged

    def _apply_block_payload_to_document(self, document: dict[str, Any], block: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any]:
        doc = json.loads(json.dumps(document, default=str))
        for key in ("title", "subtitle", "executive_summary"):
            if payload.get(key):
                doc[key] = payload[key]
        if block["name"] == "executive_summary":
            return doc
        section_map = {str(section.get("title") or ""): section for section in (doc.get("sections") or []) if section.get("title")}
        for section in payload.get("sections") or []:
            title = str(section.get("title") or "").strip()
            if title and title in section_map:
                original = section_map[title]
                original.update(section)
            elif title:
                (doc.setdefault("sections", [])).append(section)
        if payload.get("recommendations"):
            doc["recommendations"] = payload["recommendations"]
        return doc

    def _log_export_block_event(
        self,
        *,
        task_id: str,
        export_type: str,
        provider: str,
        model: str,
        block_name: str,
        action: str,
        outcome: str,
        attempt_number: int,
    ) -> None:
        log_event(
            "INFO",
            "Export block outcome",
            action=action,
            task_id=task_id,
            export_type=export_type,
            provider=provider,
            model=model,
            block_name=block_name,
            outcome=outcome,
            attempt_number=attempt_number,
        )

    def _build_export_generation_jobs(self, export_context: dict[str, Any]) -> list[dict[str, Any]]:
        requested = [str(item).lower() for item in (export_context.get("formats") or [])]
        if not requested:
            requested = ["pdf", "docx", "pptx"]
        jobs: list[dict[str, Any]] = []
        report_formats = []
        if "pdf" in requested:
            report_formats.append("pdf")
        if "word" in requested or "docx" in requested:
            report_formats.append("docx")
        report_formats = list(dict.fromkeys(report_formats))
        if report_formats:
            jobs.append(
                {
                    "job_type": "report",
                    "expected_formats": report_formats,
                    "max_tokens": 1800,
                    "retry_max_tokens": 1100,
                }
            )
        if "pptx" in requested:
            jobs.append(
                {
                    "job_type": "pptx",
                    "expected_formats": ["pptx"],
                    "max_tokens": 950,
                    "retry_max_tokens": 650,
                }
            )
        return jobs

    def _complete_fragment_for_requested_formats(self, parsed: dict[str, Any], expected_formats: list[str]) -> dict[str, Any]:
        fragment: dict[str, Any] = {"summary": str(parsed.get("summary") or "").strip()}
        report_payload = parsed.get("report")
        if isinstance(report_payload, dict):
            for fmt in expected_formats:
                if fmt in {"pdf", "docx"}:
                    fragment[fmt] = json.loads(json.dumps(report_payload))
        for fmt in expected_formats:
            payload = parsed.get(fmt)
            if isinstance(payload, dict):
                fragment[fmt] = payload
        for fmt in expected_formats:
            if fmt in {"pdf", "docx"} and fmt not in fragment:
                source = fragment.get("pdf") or fragment.get("docx")
                if isinstance(source, dict):
                    fragment[fmt] = json.loads(json.dumps(source))
        if not fragment.get("summary"):
            primary = None
            for fmt in ("pdf", "docx", "pptx"):
                payload = fragment.get(fmt)
                if isinstance(payload, dict):
                    primary = payload
                    break
            if isinstance(primary, dict):
                fragment["summary"] = str(primary.get("executive_summary") or "").strip()
        return fragment

    def _merge_export_fragments(self, fragments: list[dict[str, Any]], export_context: dict[str, Any]) -> dict[str, Any]:
        merged: dict[str, Any] = {"summary": ""}
        for fragment in fragments:
            if not merged.get("summary") and fragment.get("summary"):
                merged["summary"] = fragment["summary"]
            for key in ("docx", "pdf", "pptx"):
                payload = fragment.get(key)
                if isinstance(payload, dict):
                    merged[key] = payload
        requested = {str(item).lower() for item in (export_context.get("formats") or [])}
        if "pdf" in requested and "pdf" not in merged and isinstance(merged.get("docx"), dict):
            merged["pdf"] = json.loads(json.dumps(merged["docx"]))
        if ("word" in requested or "docx" in requested) and "docx" not in merged and isinstance(merged.get("pdf"), dict):
            merged["docx"] = json.loads(json.dumps(merged["pdf"]))
        for key in ("docx", "pdf", "pptx"):
            merged.setdefault(key, {"title": "", "subtitle": "", "executive_summary": "", "sections": [], "visuals": [], "recommendations": [], "file_manifest": [], "quality_checks": []})
        if not merged.get("summary"):
            merged["summary"] = str((merged.get("pdf") or {}).get("executive_summary") or (merged.get("docx") or {}).get("executive_summary") or (merged.get("pptx") or {}).get("executive_summary") or "").strip()
        return merged

    def _build_retry_export_context(self, export_context: dict[str, Any], *, reduced_scope: bool) -> dict[str, Any]:
        retry_context = json.loads(json.dumps(export_context or {}, default=str))
        export_options = dict(retry_context.get("export_options") or {})
        selected_sections = list(export_options.get("selected_narrative_sections") or retry_context.get("aurora_sections") or [])
        if reduced_scope:
            keep = {
                "Structural interpretation",
                "Anomaly catalogue",
                "Coverage gap analysis",
                "Drilling recommendations",
                "QA summary",
            }
            selected_sections = [item for item in selected_sections if item in keep][:2]
        retry_context["aurora_sections"] = selected_sections
        export_options["selected_narrative_sections"] = selected_sections
        retry_context["export_options"] = export_options
        return retry_context

    def _primary_failure_class(self, parse_meta: dict[str, Any]) -> str:
        for forensic in parse_meta.get("forensics") or []:
            failure_class = str(forensic.get("failure_class") or "").strip()
            if failure_class:
                return failure_class
        return ""

    def _log_export_retry_mode(
        self,
        *,
        task_id: str,
        export_type: str,
        provider: str,
        model: str,
        retry_mode: str,
        attempt_number: int,
        reason: str,
        block_name: str | None = None,
    ) -> None:
        log_event(
            "INFO",
            "Export retry mode",
            action="export.path.retry",
            task_id=task_id,
            export_type=export_type,
            provider=provider,
            model=model,
            retry_mode=retry_mode,
            attempt_number=attempt_number,
            reason=reason,
            block_name=block_name,
        )

    def _run_export_generation_job(
        self,
        *,
        project: dict,
        task: dict,
        config: dict,
        full_results: dict,
        stats: dict,
        upload_context: str | None,
        export_context: dict[str, Any],
        project_id: str,
        task_id: str,
        provider: str,
        model: str,
        job: dict[str, Any],
    ) -> tuple[dict[str, Any] | None, str, list[str]]:
        errors: list[str] = []
        export_type = ",".join(export_context.get("formats") or []) or "unspecified"
        expected_formats = list(job.get("expected_formats") or [])
        retry_mode = "same_request_retry"
        retry_export_context = dict(export_context)
        try:
            user_prompt = self._build_export_data_prompt(
                project,
                task,
                config,
                full_results,
                stats,
                upload_context,
                export_context,
                expected_formats=expected_formats,
                retry_mode="initial",
            )
        except Exception as exc:
            self._log_export_failure(
                task_id=task_id,
                export_type=export_type,
                provider=provider,
                model=model,
                attempt_number=0,
                phase="prompt_build",
                category="prompt_build_failed",
                message=str(exc),
            )
            raise
        self._log_export_token_estimate(task_id, project_id, _EXPORT_SYSTEM_PROMPT, user_prompt)
        try:
            text = self._export_client.generate(
                system_prompt=_EXPORT_SYSTEM_PROMPT,
                user_prompt=user_prompt,
                max_tokens=int(job.get("max_tokens") or 2000),
            )
            parsed, parse_meta = self._parse_export_json(text, expected_formats=expected_formats)
            if parsed:
                return self._complete_fragment_for_requested_formats(parsed, expected_formats), "success", errors
            self._log_export_parse_forensics(
                task_id=task_id,
                export_type=export_type,
                provider=provider,
                model=model,
                attempt_number=1,
                phase="api_response_parse",
                forensics=parse_meta.get("forensics") or [],
            )
            self._log_export_failure(
                task_id=task_id,
                export_type=export_type,
                provider=provider,
                model=model,
                attempt_number=1,
                phase="api_response_parse",
                category=parse_meta.get("failure_category") or "response_parse_failed",
                message="Export response could not be parsed into a valid package.",
                parse_strategy_attempted=parse_meta.get("strategies"),
                fenced_json_detected=parse_meta.get("fenced_json_detected"),
                object_extraction_attempted=parse_meta.get("object_extraction_attempted"),
                object_candidate_count=parse_meta.get("object_candidate_count"),
                response_preview=self._safe_preview(text),
            )
            failure_class = self._primary_failure_class(parse_meta)
            retry_expected_formats = list(expected_formats)
            if failure_class == "truncated_json":
                if str(job.get("job_type") or "") == "report":
                    retry_mode = "reduced_scope_retry"
                    retry_export_context = self._build_retry_export_context(export_context, reduced_scope=True)
                else:
                    retry_mode = "compact_schema_retry"
                    retry_export_context = self._build_retry_export_context(export_context, reduced_scope=False)
            self._log_export_retry_mode(
                task_id=task_id,
                export_type=export_type,
                provider=provider,
                model=model,
                retry_mode=retry_mode,
                attempt_number=2,
                reason=failure_class or "response_parse_failed",
            )
        except ExportProviderError as exc:
            errors.append(str(exc))
            self._log_export_failure(
                task_id=task_id,
                export_type=export_type,
                provider=exc.provider,
                model=exc.model,
                attempt_number=1,
                phase="api_request",
                category=exc.category,
                message=exc.provider_message,
                request_id=exc.request_id,
                status_code=exc.status_code,
                retryable=exc.retryable,
            )
            return None, "request_failed", errors
        except Exception as exc:
            self._log_export_failure(
                task_id=task_id,
                export_type=export_type,
                provider=provider,
                model=model,
                attempt_number=1,
                phase="api_request",
                category="unexpected_request_failure",
                message=str(exc),
            )
            logger.exception("Aurora export report generation failed", extra={"task_id": task_id, "project_id": project_id, "error": str(exc)})
            return None, "request_failed", errors

        try:
            compact_prompt = self._build_export_data_prompt(
                project,
                task,
                config,
                full_results,
                stats,
                (upload_context or "")[:900],
                retry_export_context,
                expected_formats=retry_expected_formats,
                retry_mode=retry_mode,
            )
            text = self._export_client.generate(
                system_prompt=_EXPORT_SYSTEM_PROMPT,
                user_prompt=compact_prompt,
                max_tokens=int(job.get("retry_max_tokens") or 1200),
            )
            parsed, parse_meta = self._parse_export_json(text, expected_formats=retry_expected_formats)
            if parsed:
                return self._complete_fragment_for_requested_formats(parsed, expected_formats), "success", errors
            self._log_export_parse_forensics(
                task_id=task_id,
                export_type=export_type,
                provider=provider,
                model=model,
                attempt_number=2,
                phase="api_response_parse",
                forensics=parse_meta.get("forensics") or [],
            )
            self._log_export_failure(
                task_id=task_id,
                export_type=export_type,
                provider=provider,
                model=model,
                attempt_number=2,
                phase="api_response_parse",
                category=parse_meta.get("failure_category") or "response_parse_failed",
                message="Export retry response could not be parsed into a valid package.",
                parse_strategy_attempted=parse_meta.get("strategies"),
                fenced_json_detected=parse_meta.get("fenced_json_detected"),
                object_extraction_attempted=parse_meta.get("object_extraction_attempted"),
                object_candidate_count=parse_meta.get("object_candidate_count"),
                response_preview=self._safe_preview(text),
            )
        except ExportProviderError as exc:
            errors.append(str(exc))
            self._log_export_failure(
                task_id=task_id,
                export_type=export_type,
                provider=exc.provider,
                model=exc.model,
                attempt_number=2,
                phase="api_request",
                category=exc.category,
                message=exc.provider_message,
                request_id=exc.request_id,
                status_code=exc.status_code,
                retryable=exc.retryable,
            )
            return None, "request_failed", errors
        except Exception as exc:
            self._log_export_failure(
                task_id=task_id,
                export_type=export_type,
                provider=provider,
                model=model,
                attempt_number=2,
                phase="api_request",
                category="unexpected_retry_failure",
                message=str(exc),
            )
            logger.exception("Aurora export retry failed", extra={"task_id": task_id, "project_id": project_id, "error": str(exc)})
            return None, "request_failed", errors
        return None, "response_invalid", errors

    def _parse_chat_response(self, text: str, location: str) -> AuroraResponse:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        summary = lines[0] if lines else text
        highlights = []
        for line in lines[1:]:
            clean = re.sub(r"^[-*•]+\s*", "", line).strip()
            if clean and len(clean) > 10:
                highlights.append(clean)
                if len(highlights) >= 6:
                    break
        return AuroraResponse(location=location, summary=summary, highlights=highlights, message=text)

    def _parse_export_json(self, text: str, *, expected_formats: list[str] | None = None) -> tuple[dict | None, dict[str, Any]]:
        parsed, meta = self._parse_json_object(text)
        if parsed is None:
            return None, meta
        normalized = self._normalize_export_payload_fragment(parsed, expected_formats or [])
        if normalized is not None:
            return normalized, meta
        meta["forensics"].append(
            {
                "parse_strategy": "shape_validation",
                "candidate_length": len(json.dumps(parsed)),
                "candidate_startswith_brace": True,
                "candidate_endswith_brace": True,
                "fenced_json_detected": meta.get("fenced_json_detected"),
                "object_candidate_count": None,
                "decode_error_type": "UnrecognizedExportShape",
                "decode_error_message": "JSON decoded successfully but did not match the expected export package shape.",
                "decode_error_line": None,
                "decode_error_column": None,
                "decode_error_pos": None,
                "failure_preview": self._safe_error_preview(json.dumps(parsed), None),
                "failure_class": "unknown_parse_failure",
            }
        )
        meta["failure_category"] = "response_parse_failed"
        return None, meta

    def _parse_json_object(self, text: str) -> tuple[dict | None, dict[str, Any]]:
        raw_text = text or ""
        stripped = raw_text.strip()
        truncation_signals = self._detect_truncation_signals(raw_text)
        meta = {
            "strategies": [],
            "fenced_json_detected": False,
            "object_extraction_attempted": False,
            "object_candidate_count": 0,
            "failure_category": "response_parse_failed",
            "forensics": [],
            "truncation_signals": truncation_signals,
        }
        direct, forensic = self._decode_json_object_candidate(stripped, strategy="direct_json", fenced_json_detected=False, truncation_signals=truncation_signals)
        meta["strategies"].append("direct_json")
        if forensic:
            meta["forensics"].append(forensic)
            meta["failure_category"] = self._response_parse_category(forensic)
        if direct is not None:
            return direct, meta
        if "```" in raw_text:
            meta["fenced_json_detected"] = True
        fenced_blocks = self._extract_fenced_code_blocks(raw_text)
        if fenced_blocks:
            meta["strategies"].append("fenced_json")
            if len(fenced_blocks) == 1:
                fenced, forensic = self._decode_json_object_candidate(
                    fenced_blocks[0],
                    strategy="fenced_json",
                    fenced_json_detected=meta["fenced_json_detected"],
                    truncation_signals=truncation_signals,
                )
                if forensic:
                    meta["forensics"].append(forensic)
                    meta["failure_category"] = self._response_parse_category(forensic)
                if fenced is not None:
                    return fenced, meta
        wrapper_stripped = self._strip_markdown_json_wrappers(raw_text)
        if wrapper_stripped is not None:
            meta["strategies"].append("wrapper_stripped_json")
            stripped_result, forensic = self._decode_json_object_candidate(
                wrapper_stripped,
                strategy="wrapper_stripped_json",
                fenced_json_detected=meta["fenced_json_detected"],
                truncation_signals=truncation_signals,
            )
            if forensic:
                meta["forensics"].append(forensic)
                meta["failure_category"] = self._response_parse_category(forensic)
            if stripped_result is not None:
                return stripped_result, meta
        meta["object_extraction_attempted"] = True
        meta["strategies"].append("object_extraction")
        candidates = self._extract_json_object_candidates(raw_text)
        meta["object_candidate_count"] = len(candidates)
        if len(candidates) != 1:
            forensic = {
                "parse_strategy": "object_extraction",
                "candidate_length": 0,
                "candidate_startswith_brace": False,
                "candidate_endswith_brace": False,
                "fenced_json_detected": meta["fenced_json_detected"],
                "object_candidate_count": len(candidates),
                "decode_error_type": "object_extraction_ambiguity",
                "decode_error_message": (
                    "Multiple top-level JSON objects detected."
                    if len(candidates) > 1
                    else "No recoverable top-level JSON object detected."
                ),
                "decode_error_line": None,
                "decode_error_column": None,
                "decode_error_pos": None,
                "failure_preview": "",
                "failure_class": "ambiguous_multiple_objects" if len(candidates) > 1 else "unknown_parse_failure",
            }
            meta["forensics"].append(forensic)
            if len(candidates) > 1:
                meta["failure_category"] = "response_parse_ambiguous"
            return None, meta
        decoded, forensic = self._decode_json_object_candidate(
            candidates[0],
            strategy="object_extraction",
            fenced_json_detected=meta["fenced_json_detected"],
            object_candidate_count=len(candidates),
            truncation_signals=truncation_signals,
        )
        if forensic:
            meta["forensics"].append(forensic)
            meta["failure_category"] = self._response_parse_category(forensic)
        return decoded, meta

    def _decode_json_object_candidate(
        self,
        candidate: str | None,
        *,
        strategy: str,
        fenced_json_detected: bool,
        object_candidate_count: int | None = None,
        truncation_signals: dict[str, Any] | None = None,
    ) -> tuple[dict | None, dict[str, Any] | None]:
        if not candidate:
            return None, None
        stripped_candidate = candidate.strip()
        try:
            data = json.loads(stripped_candidate)
        except json.JSONDecodeError as exc:
            forensic = {
                "parse_strategy": strategy,
                "candidate_length": len(stripped_candidate),
                "candidate_startswith_brace": stripped_candidate.startswith("{"),
                "candidate_endswith_brace": stripped_candidate.endswith("}"),
                "fenced_json_detected": fenced_json_detected,
                "object_candidate_count": object_candidate_count,
                "decode_error_type": exc.__class__.__name__,
                "decode_error_message": exc.msg,
                "decode_error_line": exc.lineno,
                "decode_error_column": exc.colno,
                "decode_error_pos": exc.pos,
                "failure_preview": self._safe_error_preview(stripped_candidate, exc.pos),
                "failure_class": self._classify_json_decode_error(exc, stripped_candidate, truncation_signals=truncation_signals),
            }
            return None, forensic
        except Exception as exc:
            forensic = {
                "parse_strategy": strategy,
                "candidate_length": len(stripped_candidate),
                "candidate_startswith_brace": stripped_candidate.startswith("{"),
                "candidate_endswith_brace": stripped_candidate.endswith("}"),
                "fenced_json_detected": fenced_json_detected,
                "object_candidate_count": object_candidate_count,
                "decode_error_type": exc.__class__.__name__,
                "decode_error_message": str(exc),
                "decode_error_line": None,
                "decode_error_column": None,
                "decode_error_pos": None,
                "failure_preview": self._safe_error_preview(stripped_candidate, None),
                "failure_class": "unknown_parse_failure",
            }
            return None, forensic
        if not isinstance(data, dict):
            return None, {
                "parse_strategy": strategy,
                "candidate_length": len(stripped_candidate),
                "candidate_startswith_brace": stripped_candidate.startswith("{"),
                "candidate_endswith_brace": stripped_candidate.endswith("}"),
                "fenced_json_detected": fenced_json_detected,
                "object_candidate_count": object_candidate_count,
                "decode_error_type": "UnexpectedTopLevelType",
                "decode_error_message": f"Top-level JSON value was {type(data).__name__}, not object.",
                "decode_error_line": None,
                "decode_error_column": None,
                "decode_error_pos": None,
                "failure_preview": self._safe_error_preview(stripped_candidate, None),
                "failure_class": "unknown_parse_failure",
            }
        return data, None

    def _decode_export_json_candidate(
        self,
        candidate: str | None,
        *,
        strategy: str,
        fenced_json_detected: bool,
        object_candidate_count: int | None = None,
        expected_formats: list[str] | None = None,
        truncation_signals: dict[str, Any] | None = None,
    ) -> tuple[dict | None, dict[str, Any] | None]:
        data, forensic = self._decode_json_object_candidate(
            candidate,
            strategy=strategy,
            fenced_json_detected=fenced_json_detected,
            object_candidate_count=object_candidate_count,
            truncation_signals=truncation_signals,
        )
        if forensic:
            return None, forensic
        if data is None:
            return None, None
        normalized = self._normalize_export_payload_fragment(data, expected_formats or [])
        if normalized is not None:
            return normalized, None
        if "report" in data or "executive_summary" in data:
            return self._convert_legacy_export_report(data), None
        return None, {
            "parse_strategy": strategy,
            "candidate_length": len(stripped_candidate),
            "candidate_startswith_brace": stripped_candidate.startswith("{"),
            "candidate_endswith_brace": stripped_candidate.endswith("}"),
            "fenced_json_detected": fenced_json_detected,
            "object_candidate_count": object_candidate_count,
            "decode_error_type": "UnrecognizedExportShape",
            "decode_error_message": "JSON decoded successfully but did not match the expected export package shape.",
            "decode_error_line": None,
            "decode_error_column": None,
            "decode_error_pos": None,
            "failure_preview": self._safe_error_preview(stripped_candidate, None),
            "failure_class": "unknown_parse_failure",
        }

    def _detect_truncation_signals(self, text: str) -> dict[str, Any]:
        candidate = (text or "").strip()
        quote_count = 0
        escaped = False
        for char in candidate:
            if escaped:
                escaped = False
                continue
            if char == "\\":
                escaped = True
            elif char == '"':
                quote_count += 1
        return {
            "missing_closing_brace": bool(candidate) and "{" in candidate and not candidate.endswith("}"),
            "unmatched_quotes": quote_count % 2 == 1,
            "unmatched_code_fence": candidate.count("```") % 2 == 1,
        }

    def _normalize_export_payload_fragment(self, data: dict[str, Any], expected_formats: list[str]) -> dict[str, Any] | None:
        required_formats = [fmt for fmt in expected_formats if fmt in {"pdf", "docx", "pptx"}]
        if not required_formats and "report" in data:
            required_formats = ["pdf"]
        normalized: dict[str, Any] = {"summary": str(data.get("summary") or "").strip()}
        report_payload = data.get("report")
        if isinstance(report_payload, dict):
            report_payload = self._validate_export_package({"report": report_payload}, required_formats=["report"])
            if report_payload:
                normalized["report"] = report_payload["report"]
        for key in ("docx", "pdf", "pptx"):
            payload = data.get(key)
            if not isinstance(payload, dict):
                continue
            validated = self._validate_export_package({key: payload}, required_formats=[key])
            if validated:
                normalized[key] = validated[key]
        if "report" in normalized or any(key in normalized for key in ("docx", "pdf", "pptx")):
            return normalized
        return None

    def _extract_fenced_code_blocks(self, text: str) -> list[str]:
        matches = re.findall(r"```(?:json)?\s*([\s\S]*?)\s*```", text or "", flags=re.IGNORECASE)
        return [match.strip() for match in matches if match and match.strip()]

    def _strip_markdown_json_wrappers(self, text: str) -> str | None:
        candidate = (text or "").strip()
        if not candidate.startswith("```"):
            return None
        candidate = re.sub(r"^```(?:json)?\s*", "", candidate, count=1, flags=re.IGNORECASE)
        candidate = re.sub(r"\s*```$", "", candidate, count=1, flags=re.IGNORECASE)
        return candidate.strip() or None

    def _extract_json_object_candidates(self, text: str) -> list[str]:
        candidates: list[str] = []
        start_index: int | None = None
        depth = 0
        in_string = False
        escaped = False
        for index, char in enumerate(text or ""):
            if start_index is None:
                if char == "{":
                    start_index = index
                    depth = 1
                    in_string = False
                    escaped = False
                continue
            if in_string:
                if escaped:
                    escaped = False
                elif char == "\\":
                    escaped = True
                elif char == '"':
                    in_string = False
                continue
            if char == '"':
                in_string = True
            elif char == "{":
                depth += 1
            elif char == "}":
                depth -= 1
                if depth == 0:
                    candidates.append(text[start_index : index + 1])
                    start_index = None
        return [candidate.strip() for candidate in candidates if candidate.strip()]

    def _build_chat_system_prompt(
        self,
        project: dict,
        task: dict,
        config: dict,
        outputs: dict,
        stats: dict,
        location: str,
        ui_context: dict,
    ) -> str:
        corrections = ", ".join(config.get("corrections") or []) or "none"
        add_ons = ", ".join(config.get("add_ons") or []) or "none"
        scenario = task.get("scenario") or "off"
        model = config.get("model") or "not specified"
        prompt = (
            "You are Aurora AI, a practical magnetic-survey assistant. "
            "Keep answers direct, plain-language, and grounded in the actual task data. "
            "Do not add geology interpretation unless the user explicitly asks.\n\n"
            "--- PROJECT CONTEXT ---\n"
            f"Project: {project.get('name', '')} - {project.get('context', '')}\n"
            f"Task: {task.get('name', '')} - {task.get('description', '')}\n"
            f"Platform: {task.get('platform', 'not specified')} | Scenario: {scenario} | Model: {model}\n"
            f"Corrections: {corrections}\n"
            f"Add-ons: {add_ons}\n"
        )
        if stats:
            prompt += (
                f"Magnetic stats: mean={stats.get('mean', 0):.2f} nT, std={stats.get('std', 0):.2f} nT, "
                f"min={stats.get('min', 0):.2f}, max={stats.get('max', 0):.2f}, "
                f"points={stats.get('point_count', 0)}, anomalies={stats.get('anomaly_count', 0)}\n"
            )
        validation = outputs.get("validation_summary") or {}
        base_station_count = self._resolved_base_station_count(outputs)
        if validation:
            prompt += (
                f"Validation status: {validation.get('status', 'unknown')}; "
                f"rows={validation.get('row_count', 0)}; base-station readings recognised={base_station_count}; warnings={len(validation.get('warnings') or [])}\n"
            )
        qa_report = outputs.get("qa_report") or {}
        if qa_report:
            prompt += (
                f"QA status: {qa_report.get('status', 'unknown')}; "
                f"quality={qa_report.get('quality_score', 'n/a')}; "
                f"fallback events={len(qa_report.get('fallback_events') or [])}\n"
            )
        correction_report = outputs.get("correction_report") or {}
        if correction_report:
            prompt += (
                f"Correction report: IGRF applied={bool(correction_report.get('igrf_applied'))}; "
                f"diurnal method={((correction_report.get('diurnal') or {}).get('diurnal_method') or 'not reported')}; "
                f"base readings={int(((correction_report.get('diurnal') or {}).get('n_base_readings', 0) or 0))}; "
                f"diurnal max correction={self._format_optional_float(((correction_report.get('diurnal') or {}).get('diurnal_max_correction_nT')))} nT\n"
            )
        correction_path = qa_report.get("correction_path") or []
        if correction_path:
            prompt += "Applied correction path: " + " | ".join(correction_path[:8]) + "\n"
        layers = self._list_available_layers(outputs)
        if layers:
            prompt += "Available layers: " + ", ".join(layers) + "\n"
        regional_method = outputs.get("regional_method_used")
        if regional_method:
            prompt += (
                f"Regional/residual metadata: regional method={regional_method}; "
                f"setting={json.dumps(outputs.get('regional_scale_or_degree')) if outputs.get('regional_scale_or_degree') is not None else 'not reported'}; "
                f"residual definition={outputs.get('residual_definition') or 'not reported'}\n"
            )
        sample_points = (outputs.get("points") or [])[:20]
        if sample_points:
            prompt += "Sample stations (lat, lon, magnetic nT):\n"
            for point in sample_points:
                latitude = self._to_float(point.get("latitude"))
                longitude = self._to_float(point.get("longitude"))
                magnetic = self._point_value_for_context(point)
                if latitude is None or longitude is None or magnetic is None:
                    continue
                prompt += f"  {latitude:.5f}, {longitude:.5f}, {magnetic:.2f}\n"
        uploaded_context = self._collect_uploaded_file_context(task)
        if uploaded_context:
            prompt += "\n--- UPLOADED FILES ---\n" + uploaded_context[: self._MAX_CHAT_CONTEXT_CHARS] + "\n"
        results_context = self._build_results_context(outputs)
        if results_context:
            prompt += "\n--- PROCESSED OUTPUTS ---\n" + results_context[: self._MAX_CHAT_CONTEXT_CHARS] + "\n"
        ui_summary = self._build_ui_context_summary(ui_context)
        if ui_summary:
            prompt += "\n--- UI CONTEXT ---\n" + ui_summary + "\n"
        prompt += f"\n--- CURRENT SCREEN: {location.upper()} ---\n"
        return prompt

    def _build_compact_chat_system_prompt(
        self,
        project: dict,
        task: dict,
        config: dict,
        outputs: dict,
        stats: dict,
        location: str,
        ui_context: dict,
    ) -> str:
        ui_summary = self._build_ui_context_summary(ui_context, compact=True)
        return (
            "You are Aurora AI, a practical magnetic-survey assistant. "
            "Answer directly in plain language using only the project/task/results summary below. Never expose raw field names, internal keys, or code-style identifiers.\n"
            f"Project: {project.get('name', '')}\n"
            f"Task: {task.get('name', '')}\n"
            f"Screen: {location}\n"
            f"Scenario: {task.get('scenario') or 'off'}\n"
            f"Corrections: {', '.join(config.get('corrections') or []) or 'none'}\n"
            f"Outputs: {', '.join(config.get('add_ons') or []) or 'none'}\n"
            f"Stats: mean={stats.get('mean', 0):.2f} nT, range={float(stats.get('max', 0) or 0) - float(stats.get('min', 0) or 0):.2f} nT, points={stats.get('point_count', 0)}\n"
            f"Base-station readings recognised: {self._resolved_base_station_count(outputs)}\n"
            f"Available layers: {', '.join(self._list_available_layers(outputs)) or 'Corrected Magnetic Field only'}\n"
            f"{ui_summary}"
        )

    def _build_export_data_prompt(
        self,
        project: dict,
        task: dict,
        config: dict,
        full_results: dict,
        stats: dict,
        upload_context: str | None,
        export_context: dict | None,
        *,
        expected_formats: list[str] | None = None,
        retry_mode: str = "initial",
    ) -> str:
        expected_formats = [str(item).lower() for item in (expected_formats or []) if str(item).strip()]
        data = self._build_export_context_data(
            project,
            task,
            config,
            full_results,
            stats,
            upload_context,
            export_context or {},
            expected_formats=expected_formats,
            retry_mode=retry_mode,
        )
        export_spec_excerpt = self._load_export_agent_excerpt(data, export_context or {})
        prompt_limits = self._build_export_prompt_limits(expected_formats, retry_mode)
        schema_lines = [
            "Task: build one compact JSON export package.",
            f"Requested formats: {', '.join(expected_formats) or 'report'}",
            "Return compact JSON only.",
            "Schema:",
        ]
        if any(fmt in {"pdf", "docx"} for fmt in expected_formats):
            schema_lines.append("- top keys: summary, report")
            schema_lines.append("- report keys: title, subtitle, executive_summary, sections, visuals, recommendations, quality_checks")
        if "pptx" in expected_formats:
            schema_lines.append("- top keys may also include: pptx")
            schema_lines.append("- pptx keys: title, subtitle, executive_summary, sections, visuals, recommendations, quality_checks")
        schema_lines.append("- section keys: type, title, include, bullets")
        schema_lines.append("- optional section keys only when needed: visual_refs, caption, interpretation, implication, speaker_notes, layout, callouts")
        return (
            "\n".join(schema_lines)
            + "\n\nConstraints:\n"
            + f"- retry mode: {retry_mode}\n"
            + f"- max sections: {prompt_limits['max_sections']}\n"
            + f"- max bullets per section: {prompt_limits['max_bullets']}\n"
            + f"- max bullet chars: {prompt_limits['max_bullet_chars']}\n"
            + f"- max executive summary chars: {prompt_limits['max_summary_chars']}\n"
            + f"- max interpretation chars: {prompt_limits['max_interpretation_chars']}\n"
            + f"- max caption chars: {prompt_limits['max_caption_chars']}\n"
            + f"- max speaker notes chars: {prompt_limits['max_speaker_notes_chars']}\n"
            + "- use only allowed_export_sections, selected_layers, selected_visual_refs, qa_summary, processing_path, diagnostics_summary, and findings_summary\n"
            + "- omit unsupported sections, absent layers, empty optional fields, and duplicate statements\n"
            + "- no raw dumps, placeholders, filler, repeated paragraphs, or invented geology\n"
            + "- keep QA/fallback honesty in workflow, reliability, conclusions, and recommendations\n"
            + "- return compact completed JSON only; do not emit prose, markdown fences, or comments\n"
            + "- if space is tight, shorten strings and drop optional fields before dropping required sections\n"
            + "- do not begin a string field that cannot be completed within the token limit\n"
            + "- recommendations must be limitation-driven\n"
            + "- use exact visual_refs only\n"
            + ("- PPTX: <=5 bullets per slide, no long paragraphs, one idea per slide\n" if "pptx" in expected_formats else "")
            + "\n"
            + (f"Relevant export rules:\n{export_spec_excerpt}\n\n" if export_spec_excerpt else "")
            + "Context JSON:\n"
            + json.dumps(data, separators=(",", ":"), ensure_ascii=True)
            + "\nReturn only JSON."
        )

    def _load_export_agent_excerpt(self, context_data: dict[str, Any], export_context: dict[str, Any]) -> str:
        try:
            text = _cached_export_agent_text()
            if not text:
                return ""
            formats = {str(item).lower() for item in (export_context.get("formats") or [])}
            headings = [
                "CORE PRINCIPLE",
                "GLOBAL RULES",
                "PART 1 - AI REPORTING OUTPUTS",
                "PART 7 - VISUAL INTEGRATION RULES",
                "PART 8 - TERMINOLOGY CONSISTENCY",
                "PART 9 - FALLBACK HONESTY",
            ]
            if "pptx" in formats:
                headings.append("PART 2 - PPTX-SPECIFIC RULES")
            if {"pdf", "word", "docx"} & formats:
                headings.append("PART 3 - DOCX/PDF FORMATTING RULES")
            if context_data.get("actual_results", {}).get("selected_layers"):
                headings.append("PART 6 - EXPORT SELECTION ENFORCEMENT")
            excerpts = []
            for heading in headings:
                excerpt = self._extract_markdown_heading_block(text, heading)
                if excerpt:
                    excerpts.append(excerpt)
            combined = "\n\n".join(excerpts)
            return combined[:4200]
        except Exception:
            return ""

    def _build_export_prompt_limits(self, expected_formats: list[str], retry_mode: str) -> dict[str, int]:
        is_pptx_only = expected_formats == ["pptx"]
        if retry_mode == "compact_schema_retry":
            return {
                "max_sections": 7 if is_pptx_only else 8,
                "max_bullets": 3 if is_pptx_only else 4,
                "max_bullet_chars": 90,
                "max_summary_chars": 320,
                "max_interpretation_chars": 150,
                "max_caption_chars": 90,
                "max_speaker_notes_chars": 160,
            }
        if retry_mode == "reduced_scope_retry":
            return {
                "max_sections": 7,
                "max_bullets": 4,
                "max_bullet_chars": 110,
                "max_summary_chars": 420,
                "max_interpretation_chars": 170,
                "max_caption_chars": 100,
                "max_speaker_notes_chars": 180,
            }
        return {
            "max_sections": 8 if is_pptx_only else 10,
            "max_bullets": 5 if is_pptx_only else 4,
            "max_bullet_chars": 110 if is_pptx_only else 120,
            "max_summary_chars": 420 if is_pptx_only else 520,
            "max_interpretation_chars": 180,
            "max_caption_chars": 110,
            "max_speaker_notes_chars": 220,
        }

    def _build_export_context_data(
        self,
        project: dict,
        task: dict,
        config: dict,
        full_results: dict,
        stats: dict,
        upload_context: str | None,
        export_context: dict,
        *,
        expected_formats: list[str] | None = None,
        retry_mode: str = "initial",
    ) -> dict[str, Any]:
        corrections = config.get("corrections") or []
        add_ons = config.get("add_ons") or []
        qa_report = full_results.get("qa_report") or {}
        validation = full_results.get("validation_summary") or {}
        model_meta = full_results.get("model_metadata") or {}
        processing_quality_score = qa_report.get("quality_score", 1.0)
        available_layers = self._list_available_layers(full_results)
        available_visual_assets = self._build_export_visual_refs(full_results)
        selected_sections = list(export_context.get("aurora_sections") or [])
        export_options = dict(export_context.get("export_options") or {})
        scenario = task.get("scenario") or "automatic"
        acquisition_summary = self._build_export_acquisition_summary(task, full_results, stats, validation)
        derived_findings = self._build_export_findings(task, config, full_results, stats)
        selected_layer_labels = self._selected_export_layer_labels(full_results, export_options)
        prompt_limits = self._build_export_prompt_limits(expected_formats or [], retry_mode)
        qa_summary = self._build_compact_qa_summary(full_results)
        diagnostics_summary = self._build_compact_diagnostics_summary(full_results, stats, model_meta)
        processing_path = self._build_compact_processing_path(task, full_results, corrections, add_ons)
        data = {
            "project": {
                "name": project.get("name", ""),
                "context": self._compact_text(project.get("context", ""), 220),
                "task_name": task.get("name", ""),
                "task_description": self._compact_text(task.get("description", ""), 220),
                "platform": task.get("platform", "not specified"),
                "scenario": scenario,
                "processing_mode": task.get("processing_mode", "not specified"),
            },
            "actual_results": {
                "selected_layers": selected_layer_labels,
                "selected_visual_refs": [item["visual_ref"] for item in available_visual_assets[: prompt_limits["max_sections"]]],
                "actual_emag2_reference": bool(full_results.get("emag2_reference") or full_results.get("emag2_reference_data")),
                "rtp_label": self._rtp_label(full_results),
            },
            "allowed_export_sections": self._selected_export_section_titles(full_results, export_context),
            "export_request": {
                "formats": expected_formats or export_context.get("formats") or [],
                "selected_sections": selected_sections,
                "toggles": {key: value for key, value in export_options.items() if value},
                "retry_mode": retry_mode,
            },
            "processing_path": processing_path,
            "qa_summary": qa_summary,
            "diagnostics_summary": diagnostics_summary,
            "acquisition_summary": acquisition_summary,
            "findings_summary": self._compact_findings_summary(derived_findings),
            "validation_summary": self._build_compact_validation_summary(validation),
            "residual_summary": {
                "definition": self._compact_text(full_results.get("residual_definition"), 160),
                "regional_vs_residual": self._compact_text(full_results.get("regional_vs_residual_summary"), 180),
            },
            "uploaded_files_summary": self._compact_text(upload_context[: self._MAX_UPLOAD_SNIPPET_CHARS] if upload_context else "", 280 if retry_mode != "initial" else 600),
            "prompt_limits": prompt_limits,
            "token_budget": {"optimize_for": "low_input_tokens_high_quality"},
        }
        return self._sanitize_export_context_payload(data)

    def _build_export_visual_refs(self, full_results: dict) -> list[dict[str, Any]]:
        visual_map = [
            ("Corrected Magnetic Field", "corrected_heatmap", "corrected_heatmap.png", "heatmap", "corrected_heatmap_b64", "Corrected magnetic field heatmap."),
            ("Corrected Magnetic Field", "corrected_contour", "corrected_contour.png", "contour", "corrected_contour_b64", "Corrected magnetic field contour map."),
            ("Corrected Magnetic Field", "corrected_surface", "surface.png", "surface", "surface", "Corrected magnetic field surface preview."),
            ("Regional Magnetic Field", "regional_heatmap", "regional_heatmap.png", "heatmap", "regional_heatmap_b64", "Regional magnetic field heatmap."),
            ("Regional Magnetic Field", "regional_contour", "regional_contour.png", "contour", "regional_contour_b64", "Regional magnetic field contour map."),
            ("Regional Magnetic Field", "regional_surface", "regional_surface.png", "surface", "regional_surface", "Regional magnetic field surface preview."),
            ("Residual Magnetic Field", "residual_heatmap", "residual_heatmap.png", "heatmap", "residual_heatmap_b64", "Residual magnetic field heatmap."),
            ("Residual Magnetic Field", "residual_contour", "residual_contour.png", "contour", "residual_contour_b64", "Residual magnetic field contour map."),
            ("Residual Magnetic Field", "residual_surface", "residual_surface.png", "surface", "residual_surface", "Residual magnetic field surface preview."),
            (self._rtp_label(full_results) or "Reduction to Pole (RTP)", "rtp", "rtp.png", "map", "rtp_surface", f"{self._rtp_label(full_results) or 'Reduction to Pole (RTP)'} map."),
            ("Analytic Signal", "analytic_signal", "analytic_signal.png", "map", "analytic_signal", "Analytic Signal map."),
            ("Tilt Derivative", "tilt_derivative", "tilt_derivative.png", "map", "tilt_derivative", "Tilt Derivative map."),
            ("Total Gradient", "total_gradient", "total_gradient.png", "map", "total_gradient", "Total Gradient map."),
            ("First Vertical Derivative", "first_vertical_derivative", "first_vertical_derivative.png", "map", "first_vertical_derivative", "First Vertical Derivative map."),
            ("Horizontal Derivative", "horizontal_derivative", "horizontal_derivative.png", "map", "horizontal_derivative", "Horizontal Derivative map."),
            ("Uncertainty", "uncertainty", "uncertainty.png", "map", "uncertainty", "Uncertainty map."),
            ("Line Profiles", "line_profile", "line_profile.png", "profile", "line_profile_b64", "Line profile view."),
        ]
        available: list[dict[str, Any]] = []
        for layer_label, slug, file_name, figure_type, data_key, caption in visual_map:
            if full_results.get(data_key) is None:
                continue
            available.append(
                {
                    "layer": layer_label,
                    "visual_ref": file_name,
                    "type": figure_type,
                    "slug": slug,
                    "caption": caption,
                }
            )
        return available

    def _build_export_acquisition_summary(self, task: dict, full_results: dict, stats: dict, validation: dict) -> dict[str, Any]:
        points = full_results.get("points") or []
        coordinate_system = validation.get("coordinate_system") or ((task.get("column_mapping") or {}).get("coordinate_system")) or "not reported"
        line_ids = {str(point.get("line_id")) for point in points if point.get("line_id") is not None}
        support_geometry = "single-traverse dominated" if len(line_ids) <= 1 else "multi-line support"
        return {
            "point_count": stats.get("point_count", len(points)),
            "coordinate_system": coordinate_system,
            "base_station_count": int(validation.get("base_station_count", 0) or 0),
            "line_count": validation.get("line_count", len(line_ids)),
            "station_spacing_m": self._infer_station_spacing(points),
            "support_geometry": support_geometry,
            "scenario": task.get("scenario") or "automatic",
            "platform": task.get("platform") or "not specified",
        }

    def _infer_station_spacing(self, points: list[dict]) -> float | None:
        distances = []
        ordered = [point for point in points if point.get("along_line_m") is not None and not point.get("is_base_station")]
        ordered.sort(key=lambda item: (item.get("line_id") or 0, item.get("along_line_m") or 0))
        previous_by_line: dict[str, float] = {}
        for point in ordered:
            line_key = str(point.get("line_id") or "line")
            current = self._to_float(point.get("along_line_m"))
            if current is None:
                continue
            previous = previous_by_line.get(line_key)
            if previous is not None:
                gap = current - previous
                if gap > 0:
                    distances.append(gap)
            previous_by_line[line_key] = current
        if not distances:
            return None
        distances.sort()
        return round(distances[len(distances) // 2], 2)

    def _grid_stats(self, values: Any) -> dict[str, float] | None:
        if not isinstance(values, list) or not values:
            return None
        flattened = []
        for row in values:
            if isinstance(row, list):
                flattened.extend(value for value in row if isinstance(value, (int, float)))
            elif isinstance(row, (int, float)):
                flattened.append(row)
        if not flattened:
            return None
        mean = sum(flattened) / len(flattened)
        variance = sum((value - mean) ** 2 for value in flattened) / len(flattened)
        return {"min": min(flattened), "max": max(flattened), "mean": mean, "std": variance ** 0.5, "range": max(flattened) - min(flattened)}

    def _dominant_grid_gradient(self, values: Any) -> str | None:
        try:
            import numpy as np

            arr = np.asarray(values, dtype=float)
            if arr.ndim != 2 or min(arr.shape) < 2:
                return None
            gy, gx = np.gradient(arr)
            mean_gx = float(np.nanmean(np.abs(gx)))
            mean_gy = float(np.nanmean(np.abs(gy)))
            if mean_gx <= 0 and mean_gy <= 0:
                return None
            if mean_gx > mean_gy * 1.2:
                return "east-west gradient dominance"
            if mean_gy > mean_gx * 1.2:
                return "north-south gradient dominance"
            return "mixed gradient orientation"
        except Exception:
            return None

    def _build_export_findings(self, task: dict, config: dict, full_results: dict, stats: dict) -> dict[str, Any]:
        corrected = self._grid_stats(full_results.get("surface"))
        regional = self._grid_stats(full_results.get("regional_field") or full_results.get("regional_surface"))
        residual = self._grid_stats(full_results.get("regional_residual") or full_results.get("residual_surface"))
        analytic = self._grid_stats(full_results.get("analytic_signal"))
        fvd = self._grid_stats(full_results.get("first_vertical_derivative"))
        hd = self._grid_stats(full_results.get("horizontal_derivative"))
        uncertainty = self._grid_stats(full_results.get("uncertainty"))
        findings: list[str] = []
        if corrected:
            findings.append(
                f"The corrected magnetic field spans about {corrected['range']:.2f} nT, setting the parent response against which the regional and residual products should be judged."
            )
        if regional and residual:
            relationship = "pronounced residual contrast" if residual["std"] > regional["std"] * 0.8 else "subdued residual contrast"
            findings.append(
                f"Regional separation leaves {relationship}, indicating that the run contains {'localized magnetic variation' if relationship == 'pronounced residual contrast' else 'a stronger broad background trend'} after the long-wavelength field is removed."
            )
        if analytic or fvd or hd:
            derivative_terms = []
            if analytic:
                derivative_terms.append("analytic signal")
            if fvd:
                derivative_terms.append("first vertical derivative")
            if hd:
                derivative_terms.append("horizontal derivative")
            findings.append(
                f"Derivative behaviour is expressed through {', '.join(derivative_terms)}, which helps test whether residual anomalies sharpen into clearer edges or remain diffuse."
            )
        if uncertainty:
            findings.append(
                f"Uncertainty ranges across about {uncertainty['range']:.2f} units, so confidence should remain highest close to measured support and more restrained where support thins."
            )
        corrected_orientation = self._dominant_grid_gradient(full_results.get("surface"))
        residual_orientation = self._dominant_grid_gradient(full_results.get("regional_residual") or full_results.get("residual_surface"))
        return {
            "key_findings": findings[:4],
            "corrected_orientation": corrected_orientation,
            "residual_orientation": residual_orientation,
            "observed_only_surface": not bool(config.get("run_prediction", True)),
            "regional_stats": regional,
            "residual_stats": residual,
            "corrected_stats": corrected,
        }

    def _sanitize_export_context_payload(self, data: dict[str, Any]) -> dict[str, Any]:
        payload = json.loads(json.dumps(data, default=str))
        payload.pop("sample_stations", None)
        uploaded = payload.get("uploaded_files_summary")
        if isinstance(uploaded, str):
            uploaded = re.sub(r"(?im)^\s*[\w\s-]*,?latitude,longitude,.*$", "", uploaded)
            payload["uploaded_files_summary"] = uploaded[: self._MAX_UPLOAD_SNIPPET_CHARS]
        return payload

    def _build_compact_qa_summary(self, full_results: dict) -> dict[str, Any]:
        qa_report = full_results.get("qa_report") or {}
        warnings = [self._compact_text(item, 140) for item in (qa_report.get("warnings") or []) if item][:4]
        fallbacks = []
        for event in (qa_report.get("fallback_events") or [])[:4]:
            if not isinstance(event, dict):
                continue
            fallbacks.append(
                {
                    "requested": event.get("requested"),
                    "actual": event.get("actual"),
                    "reason": self._compact_text(event.get("reason"), 140),
                    "severity": event.get("severity"),
                }
            )
        return {
            "status": qa_report.get("status", "unknown"),
            "quality_score": qa_report.get("quality_score"),
            "warnings": warnings,
            "fallback_events": fallbacks,
        }

    def _build_compact_diagnostics_summary(self, full_results: dict, stats: dict, model_meta: dict) -> dict[str, Any]:
        points = full_results.get("points") or []
        return {
            "stats": {
                "point_count": stats.get("point_count", 0),
                "range_nT": round(float(stats.get("max", 0) or 0) - float(stats.get("min", 0) or 0), 3),
                "std_nT": round(float(stats.get("std", 0) or 0), 3),
                "anomaly_count": stats.get("anomaly_count", 0),
            },
            "model": {
                "used": full_results.get("model_used") or model_meta.get("algorithm"),
                "rmse": model_meta.get("rmse", stats.get("rmse")),
                "r2": model_meta.get("r2"),
                "support_fraction": model_meta.get("support_fraction"),
                "observed_points": len([p for p in points if not p.get("is_predicted")]),
                "predicted_points": len(full_results.get("predicted_points") or []),
            },
        }

    def _build_compact_processing_path(self, task: dict, full_results: dict, corrections: list[str], add_ons: list[str]) -> dict[str, Any]:
        correction_report = full_results.get("correction_report") or {}
        applied_order = correction_report.get("applied_order") or corrections
        actual_steps = [str(step) for step in applied_order if step]
        if full_results.get("regional_field") is not None:
            actual_steps.append("regional_field_generation")
        if full_results.get("regional_residual") is not None or full_results.get("residual_surface") is not None:
            actual_steps.append("residual_field_generation")
        actual_steps.extend(str(item) for item in add_ons if item and str(item) not in actual_steps)
        deduped: list[str] = []
        for step in actual_steps:
            if step not in deduped:
                deduped.append(step)
        return {
            "steps_applied": deduped[:14],
            "model": full_results.get("model_used") or (task.get("analysis_config") or {}).get("model") or "not specified",
            "regional_method": full_results.get("regional_method_used"),
        }

    def _build_compact_validation_summary(self, validation: dict) -> dict[str, Any]:
        return {
            "status": validation.get("status", "unknown"),
            "coordinate_system": validation.get("coordinate_system") or "not reported",
            "line_count": validation.get("line_count"),
            "base_station_count": validation.get("base_station_count", 0),
            "warnings": [self._compact_text(item, 120) for item in (validation.get("warnings") or [])[:3] if item],
        }

    def _compact_findings_summary(self, derived_findings: dict[str, Any]) -> dict[str, Any]:
        return {
            "key_findings": [self._compact_text(item, 180) for item in (derived_findings.get("key_findings") or [])[:4]],
            "corrected_orientation": derived_findings.get("corrected_orientation"),
            "residual_orientation": derived_findings.get("residual_orientation"),
            "observed_only_surface": derived_findings.get("observed_only_surface"),
        }

    def _compact_text(self, text: Any, max_chars: int) -> str:
        value = re.sub(r"\s+", " ", str(text or "")).strip()
        if len(value) <= max_chars:
            return value
        return value[: max_chars - 3].rstrip() + "..."

    def _extract_markdown_heading_block(self, text: str, heading: str) -> str:
        pattern = re.compile(rf"(?ms)^#{{1,6}}\s+{re.escape(heading)}\s*\n(.*?)(?=^#{{1,6}}\s+|\Z)")
        match = pattern.search(text)
        if not match:
            return ""
        body = match.group(1).strip()
        return f"{heading}\n{body}"

    def _estimate_tokens(self, text: str) -> int:
        return max(1, (len(text or "") + 3) // 4)

    def _log_export_token_estimate(self, task_id: str, project_id: str, system_prompt: str, user_prompt: str) -> None:
        system_tokens = self._estimate_tokens(system_prompt)
        user_tokens = self._estimate_tokens(user_prompt)
        log_event(
            "INFO",
            "Export prompt token estimate",
            action="export.prompt.token_estimate",
            task_id=task_id,
            project_id=project_id,
            system_tokens_est=system_tokens,
            user_tokens_est=user_tokens,
            total_tokens_est=system_tokens + user_tokens,
        )

    def _log_export_failure(
        self,
        *,
        task_id: str,
        export_type: str,
        provider: str,
        model: str,
        attempt_number: int,
        phase: str,
        category: str,
        message: str,
        request_id: str | None = None,
        status_code: int | None = None,
        retryable: bool | None = None,
        validation_reasons: list[str] | None = None,
        response_preview: str | None = None,
        parse_strategy_attempted: list[str] | None = None,
        fenced_json_detected: bool | None = None,
        object_extraction_attempted: bool | None = None,
        object_candidate_count: int | None = None,
        block_name: str | None = None,
    ) -> None:
        log_event(
            "WARNING",
            "Export path failure",
            action="export.path.failure",
            task_id=task_id,
            export_type=export_type,
            provider=provider,
            model=model,
            attempt_number=attempt_number,
            request_id=request_id,
            status_code=status_code,
            provider_error_category=category,
            provider_error_message=self._compact_text(message, 300),
            phase=phase,
            retryable=retryable,
            validation_reasons=validation_reasons or [],
            response_preview=response_preview or "",
            parse_strategy_attempted=parse_strategy_attempted or [],
            fenced_json_detected=fenced_json_detected,
            object_extraction_attempted=object_extraction_attempted,
            object_candidate_count=object_candidate_count,
            block_name=block_name,
        )

    def _log_export_parse_forensics(
        self,
        *,
        task_id: str,
        export_type: str,
        provider: str,
        model: str,
        attempt_number: int,
        phase: str,
        forensics: list[dict[str, Any]],
        block_name: str | None = None,
    ) -> None:
        for forensic in forensics:
            log_event(
                "WARNING",
                "Export parse forensics",
                action="export.path.parse_forensics",
                task_id=task_id,
                export_type=export_type,
                provider=provider,
                model=model,
                attempt_number=attempt_number,
                phase=phase,
                parse_strategy=forensic.get("parse_strategy"),
                failure_class=forensic.get("failure_class"),
                candidate_length=forensic.get("candidate_length"),
                candidate_startswith_brace=forensic.get("candidate_startswith_brace"),
                candidate_endswith_brace=forensic.get("candidate_endswith_brace"),
                fenced_json_detected=forensic.get("fenced_json_detected"),
                object_candidate_count=forensic.get("object_candidate_count"),
                decode_error_type=forensic.get("decode_error_type"),
                decode_error_message=forensic.get("decode_error_message"),
                decode_error_line=forensic.get("decode_error_line"),
                decode_error_column=forensic.get("decode_error_column"),
                decode_error_pos=forensic.get("decode_error_pos"),
                failure_preview=forensic.get("failure_preview"),
                block_name=block_name,
            )

    def _log_export_outcome(
        self,
        *,
        task_id: str,
        export_type: str,
        provider: str,
        model: str,
        outcome: str,
    ) -> None:
        log_event(
            "INFO",
            "Export path outcome",
            action="export.path.outcome",
            task_id=task_id,
            export_type=export_type,
            provider=provider,
            model=model,
            outcome=outcome,
        )

    def _safe_preview(self, text: str | None) -> str:
        value = self._compact_text(text or "", 240)
        return re.sub(r"\s+", " ", value).strip()

    def _safe_error_preview(self, text: str, position: int | None, radius: int = 40) -> str:
        if not text:
            return ""
        if position is None:
            snippet = text[: radius * 2]
        else:
            start = max(0, position - radius)
            end = min(len(text), position + radius)
            snippet = text[start:end]
        snippet = snippet.replace("\r", "\\r").replace("\n", "\\n").replace("\t", "\\t")
        return self._compact_text(snippet, radius * 2 + 20)

    def _classify_json_decode_error(self, exc: json.JSONDecodeError, candidate: str, *, truncation_signals: dict[str, Any] | None = None) -> str:
        message = (exc.msg or "").lower()
        stripped = (candidate or "").rstrip()
        tail_distance = max(0, len(stripped) - exc.pos)
        truncation_signals = truncation_signals or {}
        if "invalid \\escape" in message or "invalid control character" in message:
            return "invalid_string_escaping"
        if "extra data" in message:
            return "trailing_non_json_content"
        if "unterminated" in message:
            return "truncated_json"
        if ("expecting value" in message or "expecting ',' delimiter" in message or "expecting property name enclosed in double quotes" in message) and tail_distance <= 3:
            return "truncated_json"
        if truncation_signals.get("missing_closing_brace") or truncation_signals.get("unmatched_quotes") or truncation_signals.get("unmatched_code_fence"):
            return "truncated_json"
        if stripped and not stripped.endswith("}") and tail_distance <= 8:
            return "truncated_json"
        return "malformed_json"

    def _response_parse_category(self, forensic: dict[str, Any]) -> str:
        failure_class = forensic.get("failure_class")
        if failure_class == "ambiguous_multiple_objects":
            return "response_parse_ambiguous"
        if failure_class in {
            "malformed_json",
            "truncated_json",
            "invalid_string_escaping",
            "trailing_non_json_content",
            "unknown_parse_failure",
        }:
            return "response_parse_failed"
        return "response_parse_failed"

    def _selected_export_layer_labels(self, full_results: dict, export_options: dict[str, Any]) -> list[str]:
        explicit_labels = [str(item).strip() for item in (export_options.get("selected_output_labels") or []) if str(item).strip()]
        if explicit_labels:
            available = set(self._list_available_layers(full_results))
            return [label for label in explicit_labels if label in available]
        allowed = []
        for label in self._list_available_layers(full_results):
            lowered = label.lower()
            if lowered.startswith("regional") and not export_options.get("include_regional_field", False):
                continue
            if lowered.startswith("residual") and not export_options.get("include_residual_field", False):
                continue
            if lowered.startswith("corrected") and not export_options.get("include_corrected_field", True):
                continue
            allowed.append(label)
        return allowed

    def _selected_export_section_titles(self, full_results: dict, export_context: dict[str, Any]) -> list[str]:
        export_options = dict(export_context.get("export_options") or {})
        selected = set(export_options.get("selected_narrative_sections") or export_context.get("aurora_sections") or [])
        base_sections = [
            "Cover",
            "Executive Summary",
            "Project Overview",
            "Survey/Data Summary",
            "Processing Workflow and QA",
            "Key Findings",
            "Main Result Interpretation",
            "Limitations and Reliability",
            "Conclusions",
        ]
        for title in self._selected_export_layer_labels(full_results, dict(export_context.get("export_options") or {})):
            base_sections.append(title)
        optional_map = {
            "Structural interpretation": "Structural / Geological Interpretation",
            "Anomaly catalogue": "Anomaly Catalogue",
            "Coverage gap analysis": "Coverage Gap Analysis",
            "Drilling recommendations": "Drilling Recommendations",
            "Correction report": "Processing Workflow and QA",
            "QA summary": "Limitations and Reliability",
        }
        for raw, normalized in optional_map.items():
            if raw in selected:
                base_sections.append(normalized)
        deduped: list[str] = []
        for title in base_sections:
            if title not in deduped:
                deduped.append(title)
        return deduped

    def _rtp_label(self, full_results: dict) -> str:
        fallback_events = (full_results.get("qa_report") or {}).get("fallback_events") or []
        if any("rtp" in str(event.get("requested") or "").lower() for event in fallback_events):
            return "Low-latitude fallback transform"
        return "Reduction to Pole (RTP)" if full_results.get("rtp_surface") is not None else ""

    def _requested_export_document_keys(self, export_context: dict[str, Any]) -> list[str]:
        formats = {str(item).lower() for item in (export_context.get("formats") or [])}
        if not formats:
            return ["docx", "pdf", "pptx"]
        keys: list[str] = []
        if "word" in formats or "docx" in formats:
            keys.append("docx")
        if "pdf" in formats:
            keys.append("pdf")
        if "pptx" in formats:
            keys.append("pptx")
        return keys or ["docx", "pdf", "pptx"]

    def _prepare_export_package(
        self,
        data: dict | None,
        full_results: dict,
        export_context: dict | None,
        *,
        task_id: str | None = None,
        export_type: str | None = None,
        provider: str | None = None,
        model: str | None = None,
        attempt_number: int | None = None,
    ) -> dict | None:
        if not data:
            return None
        prepared = self._validate_export_package(data)
        if not prepared:
            self._log_export_failure(
                task_id=task_id or "",
                export_type=export_type or "unspecified",
                provider=provider or "unknown",
                model=model or "unknown",
                attempt_number=attempt_number or 0,
                phase="export_package_normalization",
                category="malformed_schema",
                message="Export package could not be normalized into the required schema.",
            )
            return None
        prepared = self._filter_export_package_to_context(prepared, full_results, export_context or {})
        issues = self._collect_export_package_issues(prepared, full_results, export_context or {})
        if issues:
            self._log_export_failure(
                task_id=task_id or "",
                export_type=export_type or "unspecified",
                provider=provider or "unknown",
                model=model or "unknown",
                attempt_number=attempt_number or 0,
                phase="export_package_validation",
                category="validation_rejected",
                message="Export package was rejected during validation.",
                validation_reasons=issues[:20],
            )
            return None
        return prepared

    def _filter_export_package_to_context(self, data: dict, full_results: dict, export_context: dict[str, Any]) -> dict:
        available_visuals = {item.get("visual_ref") for item in self._build_export_visual_refs(full_results)}
        allowed_layers = set(self._selected_export_layer_labels(full_results, dict(export_context.get("export_options") or {})))
        allowed_sections = set(self._selected_export_section_titles(full_results, export_context))
        doc_keys = self._requested_export_document_keys(export_context)
        universal_types = {
            "cover",
            "executive_summary",
            "project_overview",
            "data_summary",
            "processing_workflow",
            "key_findings",
            "data_quality",
            "conclusions",
            "quality_check",
        }
        for key in doc_keys:
            payload = data.get(key) or {}
            cleaned_sections = []
            for section in payload.get("sections") or []:
                title = str(section.get("title") or "").strip()
                section_type = str(section.get("type") or "").strip().lower()
                if section_type not in universal_types and title and title not in allowed_sections and title not in allowed_layers:
                    continue
                section["visual_refs"] = [ref for ref in (section.get("visual_refs") or []) if ref in available_visuals]
                cleaned_sections.append(section)
            payload["sections"] = cleaned_sections
            payload["visuals"] = [visual for visual in (payload.get("visuals") or []) if isinstance(visual, dict) and visual.get("visual_ref") in available_visuals]
        return data

    def _collect_export_package_issues(self, data: dict, full_results: dict, export_context: dict[str, Any]) -> list[str]:
        issues: list[str] = []
        fallback_events = (full_results.get("qa_report") or {}).get("fallback_events") or []
        fallback_text = json.dumps(fallback_events).lower()
        actual_emag2 = bool(full_results.get("emag2_reference") or full_results.get("emag2_reference_data"))
        allowed_titles = set(self._selected_export_section_titles(full_results, export_context))
        allowed_layers = set(self._selected_export_layer_labels(full_results, dict(export_context.get("export_options") or {})))
        doc_keys = self._requested_export_document_keys(export_context)
        blocked_fragments = (
            "[insert",
            "todo",
            "tbd",
            "lorem ipsum",
            "{{",
            "}}",
            "\"type\":",
            "'type':",
            "\"properties\":",
            "undefined",
            "null",
            "none",
        )
        repeated_passages: dict[str, dict[str, str]] = {}
        for key in doc_keys:
            payload = data.get(key) or {}
            executive_summary = re.sub(r"\s+", " ", str(payload.get("executive_summary") or "").strip()).lower()
            if executive_summary:
                repeated_passages.setdefault(key, {})[executive_summary] = f"{key}:executive_summary"
            for section in payload.get("sections") or []:
                title = str(section.get("title") or "").strip()
                section_type = str(section.get("type") or "").strip().lower()
                if section_type not in {"cover", "executive_summary", "project_overview", "data_summary", "processing_workflow", "key_findings", "data_quality", "conclusions", "quality_check"}:
                    if title and title not in allowed_titles and title not in allowed_layers:
                        issues.append(f"{key}:{title}:unsupported_section")
                text_parts = [title, section.get("body") or "", section.get("interpretation") or "", section.get("implication") or "", section.get("speaker_notes") or ""] + [str(item) for item in (section.get("bullets") or [])]
                for part in text_parts:
                    normalized = re.sub(r"\s+", " ", str(part).strip())
                    lower = normalized.lower()
                    if not normalized:
                        continue
                    if any(fragment in lower for fragment in blocked_fragments):
                        issues.append(f"{key}:{title}:blocked_content")
                    if "latitude,longitude" in lower or lower.count(",") > 35:
                        issues.append(f"{key}:{title}:raw_csv_like_content")
                    if normalized.startswith("{") or normalized.startswith("["):
                        issues.append(f"{key}:{title}:raw_dump_content")
                    if len(normalized) > 80:
                        seen_title = (repeated_passages.get(key) or {}).get(lower)
                        if seen_title and seen_title != f"{key}:{title}":
                            issues.append(f"{key}:{title}:repeated_passage")
                        else:
                            repeated_passages.setdefault(key, {})[lower] = f"{key}:{title}"
                    if not actual_emag2 and "emag2" in lower:
                        issues.append(f"{key}:{title}:emag2_mislabel")
                    if "reduction to pole" in lower or re.search(r"\brtp\b", lower):
                        if "rtp" in fallback_text and "fallback" not in lower and "low-latitude" not in lower and "equator" not in lower:
                            issues.append(f"{key}:{title}:rtp_mislabel")
                    if "diurnal correction" in lower and "diurnal" in fallback_text and "fallback" not in lower and "approximate" not in lower:
                        issues.append(f"{key}:{title}:diurnal_mislabel")
                if key == "pptx":
                    bullets = section.get("bullets") or []
                    if len(bullets) > 5 or len(str(section.get("body") or "")) > 400:
                        issues.append(f"{key}:{title}:slide_density")
            recommendations = " ".join(str(item) for item in (payload.get("recommendations") or []))
            if recommendations:
                low = recommendations.lower()
                if any(marker in low for marker in ("consider follow-up", "more data may help", "additional work may be needed")):
                    issues.append(f"{key}:generic_recommendations")
                if not any(term in low for term in ("tie", "crossover", "inclination", "declination", "station", "support", "uncertainty", "fallback", "level", "prediction", "base-station")):
                    issues.append(f"{key}:untethered_recommendations")
            if (full_results.get("qa_report") or {}).get("warnings") or fallback_events:
                deck_text = " ".join(str((section.get("body") or "")) + " " + " ".join(str(item) for item in (section.get("bullets") or [])) for section in (payload.get("sections") or []))
                if not any(term in deck_text.lower() for term in ("warning", "fallback", "uncertainty", "confidence", "reliability", "caution")):
                    issues.append(f"{key}:missing_qa_reflection")
        return sorted(set(issues))

    def _validate_export_package(self, data: dict, *, required_formats: list[str] | None = None) -> dict | None:
        keys = [key for key in (required_formats or ["docx", "pdf", "pptx"]) if key in {"docx", "pdf", "pptx", "report"}]
        if not keys:
            keys = ["docx", "pdf", "pptx"]
        for key in keys:
            payload = data.get(key)
            if not isinstance(payload, dict):
                return None
            payload.setdefault("title", "")
            payload.setdefault("subtitle", "")
            payload.setdefault("executive_summary", "")
            cleaned_sections = []
            for section in (payload.get("sections") or []):
                if not isinstance(section, dict) or not section.get("include", True):
                    continue
                section.setdefault("type", "section")
                section.setdefault("title", "")
                section.setdefault("bullets", [])
                section.setdefault("body", "")
                section.setdefault("visual_refs", [])
                section.setdefault("caption", "")
                section.setdefault("interpretation", "")
                section.setdefault("implication", "")
                section.setdefault("layout", "narrative")
                section.setdefault("callouts", [])
                section.setdefault("speaker_notes", "")
                section["title"] = str(section.get("title") or "").strip()
                section["type"] = str(section.get("type") or "section").strip()
                section["body"] = str(section.get("body") or "").strip()
                section["interpretation"] = str(section.get("interpretation") or "").strip()
                section["implication"] = str(section.get("implication") or "").strip()
                section["caption"] = str(section.get("caption") or "").strip()
                section["speaker_notes"] = str(section.get("speaker_notes") or "").strip()
                section["bullets"] = [str(item).strip() for item in (section.get("bullets") or []) if str(item).strip()]
                section["visual_refs"] = [str(item).strip() for item in (section.get("visual_refs") or []) if str(item).strip()]
                section["callouts"] = [str(item).strip() for item in (section.get("callouts") or []) if str(item).strip()]
                cleaned_sections.append(section)
            payload["sections"] = cleaned_sections
            payload.setdefault("visuals", [])
            payload["visuals"] = [item for item in (payload.get("visuals") or []) if isinstance(item, dict)]
            payload["recommendations"] = [str(item).strip() for item in (payload.get("recommendations") or []) if str(item).strip()]
            payload["quality_checks"] = [str(item).strip() for item in (payload.get("quality_checks") or []) if str(item).strip()]
            payload.setdefault("file_manifest", [])
            payload["file_manifest"] = [str(item).strip() for item in (payload.get("file_manifest") or []) if str(item).strip()]
        data.setdefault("summary", (data.get("pdf") or {}).get("executive_summary") or (data.get("report") or {}).get("executive_summary") or "")
        return data

    def _convert_legacy_export_report(self, data: dict) -> dict:
        report = data.get("report") or {}
        executive_summary = data.get("executive_summary") or ""
        sections = []
        if report.get("project_overview"):
            sections.append({"type": "project_overview", "title": "Project Overview", "include": True, "bullets": [], "body": report.get("project_overview"), "visual_refs": []})
        if report.get("data_description"):
            sections.append({"type": "data_summary", "title": "Data and Survey Summary", "include": True, "bullets": [], "body": report.get("data_description"), "visual_refs": []})
        for layer in (report.get("results_interpretation") or []):
            sections.append({
                "type": "layer",
                "title": layer.get("layer", "Layer"),
                "include": True,
                "bullets": [layer.get("description", ""), layer.get("interpretation", ""), layer.get("implication", "")],
                "body": layer.get("observations", ""),
                "visual_refs": [f"{layer.get('layer', 'layer').lower().replace(' ', '_')}.png"],
            })
        package = {
            "summary": executive_summary,
            "docx": {"title": "DOCX Export", "executive_summary": executive_summary, "sections": sections, "visuals": [], "recommendations": [report.get("recommendations", "")], "file_manifest": []},
            "pdf": {"title": "PDF Export", "executive_summary": executive_summary, "sections": sections, "visuals": [], "recommendations": [report.get("recommendations", "")], "file_manifest": []},
            "pptx": {"title": "PPTX Export", "executive_summary": executive_summary, "sections": sections, "visuals": data.get("slides") or [], "recommendations": [report.get("recommendations", "")], "file_manifest": []},
        }
        return self._validate_export_package(package) or package

    def _build_export_correction_report(self, full_results: dict, corrections: list[str]) -> dict[str, Any]:
        qa_report = full_results.get("qa_report") or {}
        fallback_events = qa_report.get("fallback_events") or []
        report = {}
        for correction in corrections:
            report[correction] = "Applied or configured"
        for event in fallback_events:
            requested = event.get("requested")
            if requested:
                report[requested] = {
                    "requested": requested,
                    "actual": event.get("actual"),
                    "reason": event.get("reason"),
                    "severity": event.get("severity"),
                }
        if qa_report.get("lag"):
            report["lag"] = qa_report.get("lag")
        if qa_report.get("heading"):
            report["heading"] = qa_report.get("heading")
        if qa_report.get("crossover"):
            report["leveling"] = qa_report.get("crossover")
        return report

    def _load_full_results(self, task: dict) -> dict:
        if not self._storage:
            return {}
        artifacts = (task.get("results") or {}).get("artifacts") or []
        results_artifact = next((a for a in artifacts if a.get("file_name") == "results.json"), None)
        if not results_artifact:
            return {}
        try:
            data = self._storage.download_bytes(results_artifact["bucket"], results_artifact["object_name"])
            return json.loads(data)
        except Exception:
            return {}

    def _build_preview_outputs(self, task: dict) -> dict:
        if not self._storage:
            return {}

        survey_files = task.get("survey_files") or []
        if not survey_files:
            return {}

        mapping = task.get("column_mapping") or {}
        lat_col = mapping.get("latitude")
        lon_col = mapping.get("longitude")
        mag_col = mapping.get("magnetic_field")
        if not lat_col or not lon_col:
            return {}

        coord_sys = (mapping.get("coordinate_system") or "wgs84").lower()
        utm_zone = mapping.get("utm_zone")
        utm_hemi = mapping.get("utm_hemisphere") or "N"

        try:
            frames: list[pd.DataFrame] = []
            for file_index, artifact in enumerate(survey_files):
                csv_text = self._storage.download_text(artifact["bucket"], artifact["object_name"])
                frame = pd.read_csv(io.StringIO(csv_text))
                needed = [column for column in [lat_col, lon_col] if column and column in frame.columns]
                frame = frame.dropna(subset=needed).copy()
                frame["_raw_lat"] = pd.to_numeric(frame[lat_col], errors="coerce")
                frame["_raw_lon"] = pd.to_numeric(frame[lon_col], errors="coerce")
                frame = frame.dropna(subset=["_raw_lat", "_raw_lon"])

                if mag_col and mag_col in frame.columns:
                    frame["magnetic"] = pd.to_numeric(frame[mag_col], errors="coerce")
                else:
                    frame["magnetic"] = float("nan")

                if {"hour", "minute"}.issubset(mapping):
                    hour_col = mapping.get("hour")
                    minute_col = mapping.get("minute")
                    second_col = mapping.get("second")
                    if hour_col in frame.columns and minute_col in frame.columns:
                        sec_series = pd.to_numeric(frame.get(second_col, 0), errors="coerce").fillna(0) if second_col else 0
                        frame["time_s"] = (
                            pd.to_numeric(frame[hour_col], errors="coerce").fillna(0) * 3600.0
                            + pd.to_numeric(frame[minute_col], errors="coerce").fillna(0) * 60.0
                            + sec_series
                        )

                if coord_sys == "utm":
                    if utm_zone:
                        zone = int(utm_zone)
                        hemi = str(utm_hemi)
                    else:
                        first_e = float(frame["_raw_lon"].iloc[0])
                        first_n = float(frame["_raw_lat"].iloc[0])
                        zone, hemi = _auto_utm_zone(first_e, first_n)

                    lats: list[float | None] = []
                    lons: list[float | None] = []
                    for _, row in frame.iterrows():
                        try:
                            lat, lon = _utm_to_wgs84(float(row["_raw_lon"]), float(row["_raw_lat"]), zone, hemi)
                            lats.append(lat)
                            lons.append(lon)
                        except Exception:
                            lats.append(None)
                            lons.append(None)
                    frame["latitude"] = lats
                    frame["longitude"] = lons
                else:
                    frame["latitude"] = frame["_raw_lat"]
                    frame["longitude"] = frame["_raw_lon"]

                frame = frame.dropna(subset=["latitude", "longitude"])
                frame["is_base_station"] = _infer_base_station_mask(frame, lon_col, lat_col)
                frame["line_id"] = file_index
                frame["source_file_name"] = artifact.get("file_name") or artifact.get("object_name") or f"survey_{file_index + 1}.csv"
                frames.append(frame)

            if not frames:
                return {}

            frame = pd.concat(frames, ignore_index=True)
            points_frame = frame.head(500).copy()
            keep_cols = ["latitude", "longitude", "magnetic", "is_base_station", "line_id", "source_file_name"]
            if "time_s" in points_frame.columns:
                keep_cols.append("time_s")
            points = points_frame[keep_cols].to_dict(orient="records")

            magnetic = pd.to_numeric(frame["magnetic"], errors="coerce").dropna().tolist()
            point_count = len(magnetic)
            stats = {}
            if magnetic:
                mean_value = sum(magnetic) / point_count
                variance = sum((value - mean_value) ** 2 for value in magnetic) / point_count
                stats = {
                    "mean": mean_value,
                    "std": variance ** 0.5,
                    "min": min(magnetic),
                    "max": max(magnetic),
                    "point_count": point_count,
                    "anomaly_count": sum(1 for value in magnetic if value > mean_value + (variance ** 0.5)),
                }

            return {
                "points": points,
                "stats": stats,
                "validation_summary": {
                    "status": "preview_ready",
                    "row_count": int(len(frame)),
                    "valid_row_count": int(len(frame)),
                    "base_station_count": int(pd.to_numeric(frame["is_base_station"], errors="coerce").fillna(0).astype(int).sum()),
                    "warnings": [],
                },
            }
        except Exception:
            return {}

    def _collect_uploaded_file_context(self, task: dict) -> str | None:
        if not self._storage:
            return None
        sections = []
        seen_objects: set[str] = set()
        for artifact in list(task.get("survey_files") or []) + [task.get("basemap_file")]:
            if not artifact:
                continue
            object_name = artifact.get("object_name")
            if not object_name or object_name in seen_objects:
                continue
            seen_objects.add(object_name)
            summary = self._summarize_uploaded_artifact(artifact)
            if summary:
                sections.append(summary)
        if not sections:
            return None
        return "\n\n".join(sections)[: self._MAX_REF_CHARS]

    def _summarize_uploaded_artifact(self, artifact: dict) -> str | None:
        try:
            raw = self._storage.download_bytes(artifact["bucket"], artifact["object_name"])
        except Exception:
            return None
        file_name = artifact.get("file_name") or "uploaded file"
        content = self._extract_text_from_bytes(raw, file_name.lower())
        if not content:
            return f"File: {file_name}\nCould not extract readable text, but the file is attached to the task."
        return f"File: {file_name}\n{content[: self._MAX_UPLOAD_SNIPPET_CHARS]}"

    def _extract_text_from_bytes(self, raw: bytes, file_name: str) -> str | None:
        try:
            if file_name.endswith(".csv"):
                return self._summarize_csv(raw)
            if file_name.endswith(".txt"):
                return raw.decode("utf-8", errors="replace")
            if file_name.endswith((".json", ".geojson")):
                data = json.loads(raw.decode("utf-8", errors="replace"))
                if isinstance(data, dict) and data.get("type") == "FeatureCollection":
                    features = data.get("features") or []
                    lines = [f"GeoJSON FeatureCollection with {len(features)} features."]
                    for feature in features[:10]:
                        props = feature.get("properties") or {}
                        geom_type = (feature.get("geometry") or {}).get("type", "unknown")
                        lines.append(f"{geom_type}: {props}")
                    return "\n".join(lines)
                return json.dumps(data, indent=2)[: self._MAX_UPLOAD_SNIPPET_CHARS]
            if file_name.endswith(".kml"):
                import xml.etree.ElementTree as ET

                root = ET.fromstring(raw.decode("utf-8", errors="replace"))
                names = [el.text for el in root.iter() if el.tag.endswith("name") and el.text]
                descs = [el.text for el in root.iter() if el.tag.endswith("description") and el.text]
                lines = ["KML file contents:"]
                if names:
                    lines.append("Names: " + ", ".join(names[:20]))
                if descs:
                    lines.append("Descriptions:")
                    lines.extend(descs[:5])
                return "\n".join(lines)
            if file_name.endswith((".kmz", ".zip")):
                with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                    text_parts = []
                    for name in archive.namelist():
                        lower_name = name.lower()
                        if lower_name.endswith((".txt", ".csv", ".json", ".geojson", ".kml")):
                            extracted = self._extract_text_from_bytes(archive.read(name), lower_name)
                            if extracted:
                                text_parts.append(f"Embedded file: {name}\n{extracted[:1500]}")
                    return "\n\n".join(text_parts) if text_parts else None
            if file_name.endswith(".pdf"):
                try:
                    from pdfminer.high_level import extract_text as pdf_extract

                    return pdf_extract(io.BytesIO(raw))
                except Exception:
                    text = raw.decode("latin-1", errors="replace")
                    lines = [
                        line
                        for line in text.splitlines()
                        if len(line) > 10 and sum(32 <= ord(char) < 127 for char in line) / max(len(line), 1) > 0.8
                    ]
                    return "\n".join(lines[:80])
            if file_name.endswith(".docx"):
                from docx import Document as DocxDocument

                doc = DocxDocument(io.BytesIO(raw))
                return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())
        except Exception:
            return None
        return None

    def _summarize_csv(self, raw: bytes) -> str | None:
        text = raw.decode("utf-8", errors="replace")
        try:
            rows = list(csv.reader(io.StringIO(text)))
        except Exception:
            return None
        if not rows:
            return None
        header = rows[0]
        data_rows = rows[1:]
        if len(text) <= self._MAX_UPLOAD_SNIPPET_CHARS:
            return text

        lines = [
            f"CSV with columns: {', '.join(header[:30]) or 'unknown'}",
            f"Total rows: {len(data_rows)}",
        ]

        normalized_headers = [str(cell or "").strip().lower() for cell in header]
        longitude_index = next((idx for idx, value in enumerate(normalized_headers) if value in {"long", "long.", "longitude", "lon", "x", "easting"}), None)
        latitude_index = next((idx for idx, value in enumerate(normalized_headers) if value in {"lat", "lat.", "latitude", "y", "northing"}), None)

        base_like_rows: list[list[str]] = []
        for row in data_rows:
            row_text = " ".join((cell or "").strip().lower() for cell in row)
            if re.search(r"\b(bs|base|base station|base_station)\b", row_text):
                base_like_rows.append(row)

        if base_like_rows:
            lines.append(f"Rows containing base-station labels: {len(base_like_rows)}")
            if longitude_index is not None and latitude_index is not None:
                unique_locations = []
                seen_locations = set()
                for row in base_like_rows:
                    if longitude_index >= len(row) or latitude_index >= len(row):
                        continue
                    key = (
                        str(row[longitude_index]).strip(),
                        str(row[latitude_index]).strip(),
                    )
                    if key in seen_locations:
                        continue
                    seen_locations.add(key)
                    unique_locations.append(key)
                if unique_locations:
                    lines.append(f"Unique base-station coordinate pairs in labelled rows: {len(unique_locations)}")
                    for lon_value, lat_value in unique_locations[:20]:
                        lines.append(f"Base coordinate: lon={lon_value}, lat={lat_value}")

        def _row_preview(row: list[str]) -> str:
            return ", ".join((cell or "").strip()[:60] for cell in row[: min(len(header), 12)])

        if data_rows:
            lines.append("First rows:")
            for row in data_rows[:15]:
                lines.append(_row_preview(row))

            if len(data_rows) > 15:
                lines.append("Last rows:")
                for row in data_rows[-10:]:
                    lines.append(_row_preview(row))

        if base_like_rows:
            lines.append("Base-like rows:")
            for row in base_like_rows[:25]:
                lines.append(_row_preview(row))

        lines.append("Raw CSV excerpt:")
        lines.append(text[: min(len(text), self._MAX_UPLOAD_SNIPPET_CHARS // 2)])
        return "\n".join(lines)[: self._MAX_UPLOAD_SNIPPET_CHARS]

    def _build_results_context(self, outputs: dict) -> str:
        if not outputs:
            return ""
        lines = []
        validation = outputs.get("validation_summary") or {}
        qa_report = outputs.get("qa_report") or {}
        model_metadata = outputs.get("model_metadata") or {}
        uncertainty_metadata = outputs.get("uncertainty_metadata") or {}
        correction_report = outputs.get("correction_report") or {}
        base_station_count = self._resolved_base_station_count(outputs)
        if validation:
            lines.append(
                "Validation summary: "
                f"status={validation.get('status', 'unknown')}, "
                f"rows={validation.get('row_count', 0)}, "
                f"base-station readings={base_station_count}, "
                f"warnings={len(validation.get('warnings') or [])}"
            )
        if qa_report:
            lines.append(
                "QA report: "
                f"status={qa_report.get('status', 'unknown')}, "
                f"quality={qa_report.get('quality_score', 'n/a')}, "
                f"fallbacks={len(qa_report.get('fallback_events') or [])}"
            )
        if correction_report:
            diurnal = correction_report.get("diurnal") or {}
            lines.append(
                "Correction metadata: "
                f"igrf_applied={bool(correction_report.get('igrf_applied'))}, "
                f"diurnal_method={diurnal.get('diurnal_method', 'not reported')}, "
                f"base_readings={int(diurnal.get('n_base_readings', 0) or 0)}, "
                f"diurnal_reference={self._format_optional_float(diurnal.get('diurnal_reference_value'))}, "
                f"diurnal_max_correction_nT={self._format_optional_float(diurnal.get('diurnal_max_correction_nT'))}"
            )
        if model_metadata:
            lines.append(
                "Model metadata: "
                f"algorithm={model_metadata.get('algorithm') or 'not specified'}, "
                f"rmse={model_metadata.get('rmse')}, "
                f"mae={model_metadata.get('mae')}, "
                f"r2={model_metadata.get('r2')}"
            )
        if uncertainty_metadata:
            lines.append(
                "Uncertainty metadata: "
                f"algorithm={uncertainty_metadata.get('algorithm') or 'not specified'}, "
                f"detail={uncertainty_metadata.get('detail') or 'n/a'}"
            )
        layer_summaries = []
        for label in self._list_available_layers(outputs):
            key = self._output_key_for_layer(label)
            summary = self._summarize_numeric_grid(outputs.get(key))
            if summary:
                layer_summaries.append(f"{label}: {summary}")
        if layer_summaries:
            lines.append("Layer summaries:")
            lines.extend(layer_summaries[:10])
        return "\n".join(lines)

    def _build_ui_context_summary(self, ui_context: dict, compact: bool = False) -> str:
        if not isinstance(ui_context, dict) or not ui_context:
            return ""
        lines = []
        screen_label = ui_context.get("screen_label")
        if screen_label:
            lines.append(f"UI screen: {screen_label}")
        active_view = ui_context.get("active_view")
        if active_view:
            lines.append(f"Visualisation mode: {active_view}")
        active_layer = ui_context.get("active_layer_label")
        if active_layer:
            lines.append(f"Active layer: {active_layer}")
        selected_traverse = ui_context.get("selected_traverse_label")
        if selected_traverse:
            lines.append(f"Selected traverse: {selected_traverse}")
        if ui_context.get("is_profile_plot") is True:
            lines.append("Current view is a profile plot.")
        if ui_context.get("is_map_view") is True:
            lines.append("Current view is a map-based visual.")
        profile_line = ui_context.get("profile_line")
        if isinstance(profile_line, dict):
            start = profile_line.get("start")
            end = profile_line.get("end")
            if isinstance(start, dict) and isinstance(end, dict):
                start_lat = self._format_optional_float(start.get("latitude"), digits=5)
                start_lon = self._format_optional_float(start.get("longitude"), digits=5)
                end_lat = self._format_optional_float(end.get("latitude"), digits=5)
                end_lon = self._format_optional_float(end.get("longitude"), digits=5)
                lines.append(
                    f"Approximate line endpoints: start=({start_lat}, {start_lon}), end=({end_lat}, {end_lon})"
                )
            if profile_line.get("length_m") is not None:
                lines.append(f"Approximate displayed line length: {self._format_optional_float(profile_line.get('length_m'))} m")
        displayed_stats = ui_context.get("displayed_stats")
        if isinstance(displayed_stats, dict):
            lines.append(
                "Displayed values: "
                f"min={self._format_optional_float(displayed_stats.get('min'))}, "
                f"max={self._format_optional_float(displayed_stats.get('max'))}, "
                f"mean={self._format_optional_float(displayed_stats.get('mean'))}, "
                f"range={self._format_optional_float(displayed_stats.get('range'))}"
            )
        value_source = ui_context.get("value_source")
        if value_source:
            lines.append(f"Displayed-value source: {value_source}")
        if compact:
            return ("UI: " + " | ".join(lines) + "\n") if lines else ""
        return "\n".join(lines)

    def _format_optional_float(self, value: Any, digits: int = 2) -> str:
        number = self._to_float(value)
        if number is None:
            return "n/a"
        return f"{number:.{digits}f}"

    def _resolved_base_station_count(self, outputs: dict) -> int:
        points = outputs.get("points") or []
        counted = sum(1 for point in points if point.get("is_base_station"))
        if counted:
            return int(counted)
        validation = outputs.get("validation_summary") or {}
        return int(validation.get("base_station_count", 0) or 0)

    def _list_available_layers(self, outputs: dict) -> list[str]:
        layer_map = {
            "surface": "Corrected Magnetic Field",
            "corrected_field": "Corrected Magnetic Field",
            "filtered_surface": "Filtered Surface",
            "rtp_surface": self._rtp_label(outputs) or "Reduction to Pole (RTP)",
            "analytic_signal": "Analytic Signal",
            "first_vertical_derivative": "First Vertical Derivative (FVD)",
            "horizontal_derivative": "Horizontal Derivative",
            "regional_field": "Regional Magnetic Field",
            "regional_surface": "Regional Magnetic Field",
            "regional_residual": "Residual Magnetic Field",
            "residual_surface": "Residual Magnetic Field",
            "emag2_residual": "Residual Magnetic Field",
            "uncertainty": "Uncertainty Surface",
            "tilt_derivative": "Tilt Derivative",
            "total_gradient": "Total Gradient",
        }
        labels = []
        for key, label in layer_map.items():
            if outputs.get(key) is not None and label not in labels:
                labels.append(label)
        return labels

    def _output_key_for_layer(self, label: str) -> str:
        reverse = {
            "Corrected Magnetic Field": "corrected_field",
            "Filtered Surface": "filtered_surface",
            "Reduction to Pole (RTP)": "rtp_surface",
            "Low-latitude fallback transform": "rtp_surface",
            "Analytic Signal": "analytic_signal",
            "First Vertical Derivative (FVD)": "first_vertical_derivative",
            "Horizontal Derivative": "horizontal_derivative",
            "Regional Magnetic Field": "regional_field",
            "Residual Magnetic Field": "regional_residual",
            "Uncertainty Surface": "uncertainty",
            "Tilt Derivative": "tilt_derivative",
            "Total Gradient": "total_gradient",
        }
        return reverse.get(label, "")

    def _build_fallback_export_package(
        self,
        project: dict,
        task: dict,
        config: dict,
        full_results: dict,
        stats: dict,
        upload_context: str | None,
        export_context: dict | None,
    ) -> dict:
        qa_report = full_results.get("qa_report") or {}
        validation = full_results.get("validation_summary") or {}
        layers = self._list_available_layers(full_results)
        corrections = ", ".join(config.get("corrections") or []) or "None"
        add_ons = ", ".join(config.get("add_ons") or []) or "None"
        fallback_events = qa_report.get("fallback_events") or []
        warnings = qa_report.get("warnings") or []
        provider_errors = list((export_context or {}).get("provider_errors") or [])
        quality_score = qa_report.get("quality_score")
        confidence_phrase = "moderate confidence"
        if qa_report.get("status") == "valid" and not warnings and not fallback_events:
            confidence_phrase = "relatively strong confidence"
        elif qa_report.get("status") in {"degraded", "failed"} or fallback_events:
            confidence_phrase = "cautious confidence"
        run_prediction = bool(config.get("run_prediction", True))
        observed_only = not run_prediction or str(task.get("scenario") or "").lower() in {"off", "none"}
        leveling_limited = ("level" in corrections or "micro_leveling" in corrections) and any((event.get("requested") or "").lower().startswith("level") for event in fallback_events)

        recommendations = []
        if any((event.get("requested") or "").lower().startswith("level") for event in fallback_events):
            recommendations.append("Acquire tie lines or crossover support before relying on subtle line-to-line magnetic differences.")
        if any((event.get("requested") or "").lower().startswith("rtp") for event in fallback_events):
            recommendations.append("Provide inclination and declination so RTP can be recomputed as a physically reliable transform.")
        if any((event.get("requested") or "").lower().startswith("diurnal") for event in fallback_events) or int(validation.get("base_station_count", 0) or 0) == 0:
            recommendations.append("Improve base-station coverage or timing overlap so diurnal correction is constrained by measured drift rather than fallback behavior.")
        if (task.get("scenario") or "").lower() == "sparse":
            recommendations.append("Increase station density in lower-support zones before treating short-wavelength residual features as drilling-scale targets.")
        if observed_only:
            recommendations.append("Do not project anomaly confidence beyond measured support because this run is based on observed-data representation rather than an interpolative prediction surface.")
        if not recommendations:
            recommendations = [
                "Carry the corrected, regional, residual, and derivative layers together into follow-up interpretation so anomaly scale is not judged from a single map alone.",
                "Prioritise additional field control where uncertainty or sparse support weakens confidence in local anomaly geometry.",
            ]
        if provider_errors:
            recommendations.append("Review the export AI provider health before relying on richer narrative formatting, because this package was regenerated through the deterministic fallback path.")

        executive_summary = (
            f"{task.get('name', 'Task')} for project {project.get('name', 'Project')} processed "
            f"{stats.get('point_count', 0)} points with {qa_report.get('status', 'unknown')} QA status. "
            f"The interpretation carries {confidence_phrase} after considering the correction path, support coverage, and QA outcomes."
        )
        if observed_only:
            executive_summary += " The mapped surface should be treated as an observed-data representation rather than as a fully interpolated prediction surface."
        selected_sections = list((export_context or {}).get("aurora_sections") or [])
        visuals = self._build_export_visual_refs(full_results)
        acquisition_summary = self._build_export_acquisition_summary(task, full_results, stats, validation)
        derived_findings = self._build_export_findings(task, config, full_results, stats)
        layer_sections = []
        layer_logic = {
            "Corrected Magnetic Field": {
                "bullets": [
                    "This layer anchors the overall magnetic response before regional separation.",
                    "Its value is strongest when read together with the regional and residual products rather than in isolation.",
                ],
                "body": "The corrected magnetic field establishes the full processed signal, including both broad background behaviour and shorter-wavelength anomaly content. It therefore provides the baseline against which the regional field and residual field should be compared.",
                "interpretation": "Broad amplitude changes that persist into the regional field indicate long-wavelength background structure, while departures that intensify in the residual map mark more localised magnetic sources.",
                "implication": "Use this layer to frame the total signal first, then use regional/residual separation and derivatives to decide which responses remain important at the target scale.",
            },
            "Regional Magnetic Field": {
                "bullets": [
                    "This surface isolates the longer-wavelength magnetic trend that also influences the corrected field.",
                    "Regional behaviour should be compared directly with the residual layer to separate broad trend from anomaly expression.",
                ],
                "body": "The regional field captures the broader magnetic background that underlies the corrected response. Where the corrected field varies smoothly and the residual remains subdued, the dominant behaviour is more likely regional than target-scale.",
                "interpretation": "Longer-wavelength gradients here explain the background tilt of the corrected field and help prevent those broader components from being misread as discrete local anomalies.",
                "implication": "Targets should be prioritised where residual anomalies depart clearly from this regional background rather than where the response is driven mainly by the broad trend itself.",
            },
            "Residual Magnetic Field": {
                "bullets": [
                    "Residual behaviour should be judged against both the corrected field and the regional background.",
                    "Sharper features that persist into derivatives are more likely to represent localised magnetic contrasts.",
                ],
                "body": "The residual field subtracts the regional background from the corrected field so that shorter-wavelength anomaly content becomes easier to isolate. This is the key bridge between the total field and the derivative products.",
                "interpretation": "Residual highs and lows that align with stronger edge responses in the derivatives are more consistent with local magnetic sources than with broad background drift.",
                "implication": "Use the residual layer as the main anomaly-screening surface, but confirm edge geometry and continuity with derivative layers before drawing structural or drilling conclusions.",
            },
            "Analytic Signal": {
                "bullets": [
                    "Analytic Signal should be read against the residual field rather than against the regional field alone.",
                    "It strengthens source-edge and amplitude focus where residual anomalies are already evident.",
                ],
                "body": "Analytic Signal sharpens the response around magnetic source contrasts without preserving anomaly sign in the same way as the corrected or residual fields. It therefore serves best as an edge-and-source-strength companion to the residual surface.",
                "interpretation": "Where Analytic Signal reinforces compact residual anomalies, the anomaly edges are more likely to be coherent and geologically meaningful than where only the total field is elevated.",
                "implication": "Use this layer to refine anomaly footprint and edge continuity after the residual field has isolated local magnetic behaviour.",
            },
            "First Vertical Derivative (FVD)": {
                "bullets": [
                    "FVD should be compared with the residual and horizontal derivative to judge edge sharpness rather than amplitude alone.",
                    "Sharper derivative responses are more sensitive to noise and QA limitations than the corrected field.",
                ],
                "body": "The first vertical derivative enhances short-wavelength changes and suppresses broader background behaviour, making it useful for assessing the sharpness of residual anomalies and the likelihood of near-surface magnetic contrasts.",
                "interpretation": "If FVD strengthens the same features that appear in the residual map, those responses are more likely to reflect sharper source boundaries than simple background trend.",
                "implication": "Treat this layer as an edge-sharpening diagnostic rather than a standalone target map, especially where QA warnings or support limitations reduce confidence.",
            },
            "Horizontal Derivative": {
                "bullets": [
                    "Horizontal Derivative should be compared with FVD and the residual field to judge lateral edge continuity.",
                    "Edge alignment matters more than isolated amplitude spikes in this layer.",
                ],
                "body": "The horizontal derivative emphasises lateral magnetic gradients and is most useful where it outlines the margins of anomalies already expressed in the residual or corrected field.",
                "interpretation": "Consistent edge trends here strengthen the case for laterally coherent boundaries, whereas diffuse or fragmented responses imply weaker structural resolution.",
                "implication": "Use this layer to test whether residual anomalies have laterally consistent margins before assigning structural significance.",
            },
        }
        for layer in layers[:8]:
            match_key = "Residual Magnetic Field" if layer == "Residual Magnetic Field" else layer
            logic = layer_logic.get(match_key, None)
            visual_refs = [item["visual_ref"] for item in visuals if layer.lower().split(" (")[0] in item["layer"].lower()][:2]
            caution = []
            if warnings or fallback_events:
                qa_reason = f"QA reported {len(warnings)} warning(s) and {len(fallback_events)} fallback event(s)"
                if layer == "Corrected Magnetic Field":
                    caution.append(f"The parent magnetic response is read with {confidence_phrase} because {qa_reason}.")
                elif layer == "Regional Magnetic Field":
                    caution.append(f"The regional trend should be treated with {confidence_phrase} because {qa_reason}.")
                elif layer == "Residual Magnetic Field":
                    caution.append(f"Residual anomaly contrast is read with {confidence_phrase} because {qa_reason}.")
                elif layer == "Analytic Signal":
                    caution.append(f"Edge-strength interpretation stays at {confidence_phrase} because {qa_reason}.")
                elif layer == "First Vertical Derivative (FVD)":
                    caution.append(f"Derivative sharpening carries {confidence_phrase} because {qa_reason}.")
                elif layer == "Horizontal Derivative":
                    caution.append(f"Lateral boundary mapping remains at {confidence_phrase} because {qa_reason}.")
            if any((event.get("requested") or "").lower().startswith("rtp") for event in fallback_events) and "RTP" in layer:
                caution.append("RTP was degraded by a fallback path, so map geometry here should not be treated as a physically robust pole-reduced response.")
            if leveling_limited and layer in {"Corrected Magnetic Field", "Residual Magnetic Field", "First Vertical Derivative (FVD)", "Horizontal Derivative"}:
                caution.append("Leveling confidence is limited, so subtle line-parallel bias may still influence weak gradients or low-amplitude residual features.")
            if observed_only and layer in {"Corrected Magnetic Field", "Residual Magnetic Field", "Regional Magnetic Field"}:
                if layer == "Corrected Magnetic Field":
                    caution.append("The corrected field represents measured support rather than an interpolated prediction surface beyond the observed stations.")
                elif layer == "Regional Magnetic Field":
                    caution.append("The regional trend should not be treated as a confident interpolated background away from measured support because this run remained observed-only.")
                elif layer == "Residual Magnetic Field":
                    caution.append("Residual amplitudes outside measured support should not be treated as interpolatively verified because this run stayed in observed-only mode.")
            layer_sections.append(
                {
                    "type": "layer",
                    "title": layer,
                    "include": True,
                    "bullets": (logic or {}).get("bullets", [f"{layer} should be interpreted in relation to the corrected, residual, and derivative layers rather than as a standalone map."]) + caution,
                    "body": (logic or {}).get("body", f"{layer} contributes part of the processed magnetic interpretation chain and is most useful when cross-checked against adjacent products."),
                    "visual_refs": visual_refs,
                    "figure_title": f"{layer} Interpretation",
                    "caption": f"{layer} viewed in the context of the processed magnetic workflow.",
                    "interpretation": (logic or {}).get("interpretation", f"{layer} should be read against the neighbouring layers so that broad trend, local anomaly content, and edge sharpening are not conflated."),
                    "implication": (logic or {}).get("implication", "Use this layer as supporting evidence within the full processing sequence, not as a standalone conclusion."),
                    "callouts": caution,
                    "table_rows": [],
                    "speaker_notes": (logic or {}).get("body", ""),
                }
            )
        common_sections = [
            {
                "type": "project_overview",
                "title": "Project Overview",
                "include": True,
                "bullets": [project.get("context", "")],
                "body": f"Project: {project.get('name', '')}. Task: {task.get('name', '')}.",
                "visual_refs": [],
            },
            {
                "type": "data_summary",
                "title": "Data and Survey Summary",
                "include": True,
                "bullets": [
                    f"Scenario: {acquisition_summary.get('scenario')}",
                    f"Platform: {acquisition_summary.get('platform')}",
                    f"Coordinate system: {acquisition_summary.get('coordinate_system')}",
                    f"Base-station readings recognised: {acquisition_summary.get('base_station_count')}",
                ],
                "body": (
                    f"The run is based on approximately {acquisition_summary.get('point_count', 0)} points across "
                    f"{acquisition_summary.get('line_count', 0)} line grouping(s), with "
                    f"{acquisition_summary.get('support_geometry')} support. "
                    + (
                        f"Inferred station spacing is about {acquisition_summary.get('station_spacing_m')} m where along-line support is reliable. "
                        if acquisition_summary.get("station_spacing_m") is not None
                        else "Station spacing could not be inferred reliably from the stored run outputs. "
                    )
                    + ("Raw uploaded rows are not repeated here; acquisition detail is summarised so the interpretation remains readable." if upload_context else "")
                ),
                "visual_refs": [],
                "table_rows": [
                    {"Metric": "Point count", "Value": acquisition_summary.get("point_count", 0)},
                    {"Metric": "Coordinate system", "Value": acquisition_summary.get("coordinate_system")},
                    {"Metric": "Base-station count", "Value": acquisition_summary.get("base_station_count")},
                    {"Metric": "Station spacing (m)", "Value": acquisition_summary.get("station_spacing_m") or "not inferred"},
                    {"Metric": "Support geometry", "Value": acquisition_summary.get("support_geometry")},
                ],
            },
            {
                "type": "key_findings",
                "title": "Key Findings",
                "include": True,
                "bullets": derived_findings.get("key_findings") or [
                    "The corrected, regional, residual, and derivative products should be read together to define scale, edge sharpness, and support limitations."
                ],
                "body": "These findings summarise the strongest run-specific magnetic signals before the detailed layer-by-layer discussion.",
                "visual_refs": [item["visual_ref"] for item in visuals if item["visual_ref"] in {"corrected_heatmap.png", "regional_heatmap.png", "residual_heatmap.png"}][:2],
            },
            {
                "type": "processing_workflow",
                "title": "Processing Workflow Summary",
                "include": True,
                "bullets": [
                    f"Corrections: {corrections}",
                    f"Add-ons: {add_ons}",
                    f"Model: {config.get('model') or 'not specified'}",
                ],
                "body": "Processing interpretation is conditioned by the actual correction path applied. Where a requested correction downgraded to a fallback, the resulting layers should be read with more caution than a fully supported physical workflow.",
                "visual_refs": [],
                "callouts": [f"{event.get('requested')}: {event.get('reason')}" for event in fallback_events[:3] if event.get("requested") and event.get("reason")] + provider_errors[:2],
                "table_rows": [{"Correction": key, "Status": ("fallback affected" if any((event.get("requested") or "") == key for event in fallback_events) else "configured")} for key in (config.get("corrections") or [])[:6]],
            },
            *layer_sections,
            {
                "type": "data_quality",
                "title": "Data Quality and Reliability",
                "include": True,
                "bullets": [
                    f"QA status: {qa_report.get('status', 'unknown')}",
                    f"Processing quality score: {qa_report.get('quality_score', 'n/a')}",
                    f"Base-station support recognised in QA: {validation.get('base_station_count', 0)}",
                ],
                "body": f"The overall interpretation carries {confidence_phrase}. Reliability is strongest where corrected, residual, and derivative behaviour agree, and weakens where QA warnings, sparse support, or fallback corrections affect the processing path.",
                "visual_refs": [item["visual_ref"] for item in visuals if "uncertainty" in item["visual_ref"]],
                "callouts": list(warnings[:3]) + [event.get("reason") for event in fallback_events[:3] if event.get("reason")],
                "table_rows": [
                    {"Metric": "QA status", "Value": qa_report.get("status", "unknown")},
                    {"Metric": "Quality score", "Value": quality_score if quality_score is not None else "n/a"},
                    {"Metric": "Fallback events", "Value": len(fallback_events)},
                    {"Metric": "Warnings", "Value": len(warnings)},
                ],
            },
        ]
        if "Anomaly catalogue" in selected_sections:
            common_sections.append({
                "type": "anomaly_catalogue",
                "title": "Anomaly Catalogue",
                "include": True,
                "bullets": [
                    "Formal anomaly extraction was not run as a separate catalogue stage for this export.",
                    "Observable anomaly behaviour should therefore be described qualitatively from the residual and derivative layers rather than as a ranked anomaly list.",
                ],
                "body": "Where residual amplitudes tighten into sharper derivative responses, anomaly behaviour is still evident even though a formal anomaly table was not generated.",
                "visual_refs": [item["visual_ref"] for item in visuals if item["visual_ref"] in {"residual_heatmap.png", "analytic_signal.png", "line_profile.png"}],
                "figure_title": "Qualitative Anomaly Behaviour",
                "caption": "Residual and derivative products used for qualitative anomaly screening.",
                "interpretation": "Shorter-wavelength residual features that persist into the derivatives deserve more attention than broad responses seen only in the corrected field.",
                "implication": "A formal anomaly extraction pass would be the next step if ranked target zones are required.",
            })
        if "Coverage gap analysis" in selected_sections:
            common_sections.append({
                "type": "coverage_gap_analysis",
                "title": "Coverage Gap Analysis",
                "include": True,
                "bullets": [
                    "Coverage confidence depends on station support, prediction mode, and uncertainty behaviour.",
                    "Areas that require interpolation or sit farther from measured control carry weaker constraint than densely sampled zones.",
                ],
                "body": "Where uncertainty broadens or support thins, local anomaly geometry should be treated more cautiously than in densely controlled parts of the survey.",
                "visual_refs": [item["visual_ref"] for item in visuals if "uncertainty" in item["visual_ref"]],
            })
        if "Drilling recommendations" in selected_sections:
            common_sections.append({"type": "drilling_recommendations", "title": "Drilling Recommendations", "include": True, "bullets": recommendations[:4], "body": "Recommendations are tied directly to the current processing limitations and interpretation confidence rather than to generic drilling language.", "visual_refs": []})
        if "Structural interpretation" in selected_sections:
            structural_body = "Structural resolution is limited where anomaly edges are diffuse, support is sparse, or QA fallbacks weaken confidence in derivative geometry."
            if any(label in layers for label in ["Residual Magnetic Field", "Analytic Signal", "Horizontal Derivative", "First Vertical Derivative (FVD)"]):
                structural_body = "Structural interpretation should focus on features that remain coherent from the residual field into the derivative products; where that continuity breaks down, structural confidence is limited by data geometry or QA constraints."
            common_sections.append({"type": "structural_interpretation", "title": "Structural / Geological Interpretation", "include": True, "bullets": ["Structural interpretation is strongest only where residual and derivative responses reinforce the same boundaries."], "body": structural_body, "visual_refs": [item["visual_ref"] for item in visuals if item["visual_ref"] in {"residual_contour.png", "analytic_signal.png", "horizontal_derivative.png", "first_vertical_derivative.png"}]})

        common_sections.append(
            {
                "type": "conclusions",
                "title": "Conclusions",
                "include": True,
                "bullets": [
                    "The corrected field frames the total processed magnetic response.",
                    "Regional separation defines the long-wavelength background, while the residual and derivative layers isolate the more local anomaly behaviour.",
                    f"The final interpretation carries {confidence_phrase} after considering QA, support coverage, and fallback behavior.",
                ],
                "body": "The most reliable conclusions are the ones supported consistently across the corrected field, regional trend separation, residual response, and derivative sharpening rather than by any single layer alone.",
                "visual_refs": [],
                "callouts": (["Observed-only surface logic limits confidence outside measured support."] if observed_only else []) + (["Residual line-based bias may persist because leveling support was limited."] if leveling_limited else []),
            }
        )
        package = {
            "summary": executive_summary,
            "docx": {
                "title": f"{project.get('name', 'Project')} Technical Report",
                "subtitle": "Integrated magnetic interpretation report",
                "executive_summary": executive_summary,
                "sections": common_sections,
                "visuals": visuals,
                "recommendations": recommendations,
                "file_manifest": [],
            },
            "pdf": {
                "title": f"{project.get('name', 'Project')} Executive Report",
                "subtitle": "Condensed client-facing interpretation",
                "executive_summary": executive_summary,
                "sections": common_sections,
                "visuals": visuals,
                "recommendations": recommendations,
                "file_manifest": [],
            },
            "pptx": {
                "title": f"{project.get('name', 'Project')} Presentation",
                "subtitle": "Key magnetic interpretation findings",
                "executive_summary": executive_summary,
                "sections": common_sections,
                "visuals": visuals,
                "recommendations": recommendations,
                "file_manifest": [],
            },
        }
        return self._validate_export_package(package) or package

    def _summarize_numeric_grid(self, values: Any) -> str | None:
        if not isinstance(values, list) or not values:
            return None
        flattened = []
        for row in values:
            if isinstance(row, list):
                flattened.extend(value for value in row if isinstance(value, (int, float)))
            elif isinstance(row, (int, float)):
                flattened.append(row)
        if not flattened:
            return None
        return (
            f"min={min(flattened):.2f}, max={max(flattened):.2f}, "
            f"mean={sum(flattened) / len(flattened):.2f}, cells={len(flattened)}"
        )

    def _point_value_for_context(self, point: dict) -> float | None:
        for key in ("magnetic", "corrected_magnetic", "predicted_magnetic", "value"):
            value = self._to_float(point.get(key))
            if value is not None:
                return value
        return None

    def _to_float(self, value: Any) -> float | None:
        try:
            return float(value)
        except (TypeError, ValueError):
            return None

    def _fallback_text(self, project: dict, task: dict, outputs: dict, location: str, question: str | None) -> str:
        layer_names = ", ".join(self._list_available_layers(outputs)[:4]) or "saved outputs"
        if location == "export":
            return (
                "Aurora export drafting is temporarily unavailable, so the document package was generated from the saved task data only. "
                f"Available outputs for this task include {layer_names}."
            )
        question_text = f"Your message was: {question}" if question else "Try asking again in a moment."
        return (
            "Aurora is temporarily unavailable, but your saved task and processing results are still loaded. "
            f"The current screen has access to {layer_names}. {question_text}"
        )
